import json
import tempfile
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient

import app.services.validation_service as validation_service
from app.api.routes import audit_service, auth_service, job_service
from app.core.audit import AuditEventType
from app.core.llm_cache import LLM_FORCE_REFRESH_PARAM
from app.core.tenant_config import OperatorRole
from app.core.tenant_loader import load_tenant_config
from app.main import app
from app.services.tenant_admin_service import TenantAdminService
from app.services.validation_service import run_validation_job

client = TestClient(app)

DEFAULT_API_KEY = "default-local-test-key"
DEFAULT_API_KEY_ID = "default-local"
DEFAULT_OPERATOR_PASSWORD = "Validador@2026!"
REDESIM_API_KEY = "redesim-local-test-key"


def setup_function():
    job_service._jobs.clear()
    audit_service.clear()
    auth_service.clear()


def auth_headers(api_key: str = DEFAULT_API_KEY) -> dict[str, str]:
    return {"X-API-Key": api_key}


def enable_initial_setup_storage(monkeypatch, tmp_path) -> Path:
    operator_storage_path = tmp_path / "operators.json"
    monkeypatch.setattr(auth_service, "_operator_storage_path", operator_storage_path)
    monkeypatch.setattr(auth_service, "_sqlite_store", None)
    auth_service._operator_records.clear()
    return operator_storage_path


def enable_tenant_admin_storage(monkeypatch, tmp_path) -> Path:
    tenant_storage_path = tmp_path / "runtime_tenants.json"
    monkeypatch.setenv("VALIDATOR_TENANT_STORE_PATH", str(tenant_storage_path))

    import app.api.routes as routes

    monkeypatch.setattr(
        routes,
        "tenant_admin_service",
        TenantAdminService(
            storage_path=tenant_storage_path,
            audit_service=audit_service,
        ),
    )
    return tenant_storage_path


def login_headers(
    *,
    tenant_id: str = "default",
    username: str = "default.operator",
    password: str = DEFAULT_OPERATOR_PASSWORD,
) -> tuple[dict[str, str], dict]:
    response = client.post(
        "/login",
        json={
            "tenant_id": tenant_id,
            "username": username,
            "password": password,
        },
    )
    assert response.status_code == 200
    payload = response.json()
    return auth_headers(payload["x_api_key"]), payload


def platform_admin_headers() -> dict[str, str]:
    auth_service.create_operator(
        tenant_id="default",
        username="admin.plataforma",
        password="AdminGlobal@2026",
        role=OperatorRole.PLATFORM_ADMIN,
    )
    headers, _payload = login_headers(
        username="admin.plataforma",
        password="AdminGlobal@2026",
    )
    return headers


CSV_CONTENT = (
    "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
    "001,PA-100,Mesa,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe completo,Obs\n"
    "002,,Cadeira,MarcaZ,ModeloW,SN2,Sala2,CC2,,\n"
)

DUPLICATE_SCOPE_CSV_CONTENT = (
    "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
    "001,PA-100,Mesa,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe completo,Obs\n"
    "002,,Cadeira,MarcaZ,ModeloW,SN2,Sala2,CC2,,\n"
    "001,,Armario,MarcaA,ModeloB,SN3,Sala3,CC3,Detalhe armario completo validado,\n"
)

REDESIM_CSV_CONTENT = (
    "especie_id;base_id;;item_anterior;item;descricao;marca;modelo;ns;complemento;observacao;cc;cc_descricao;local;latitude;longitude;gps;usuario;foto_complementar_memento;\n"
    "1;144;uuid-1;;001;MONITOR;Dell;P2419H;SN1;;;8327;A27;MATRIZ;;;;Leticia;;\n"
)


def test_health():
    response = client.get("/health")
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "ok"
    assert [check["name"] for check in payload["checks"]] == [
        "uploads",
        "results",
        "job_store",
    ]
    assert all(check["ok"] is True for check in payload["checks"])


def test_root_returns_api_status_when_frontend_url_is_not_configured(monkeypatch):
    monkeypatch.delenv("VALIDATOR_FRONTEND_URL", raising=False)

    response = client.get("/")

    assert response.status_code == 200
    assert response.json() == {
        "service": "validador-final-api",
        "status": "ok",
        "ui": "Configure VALIDATOR_FRONTEND_URL to redirect operators to the frontend.",
        "health": "/health",
        "docs": "/docs",
    }


def test_root_redirects_to_configured_frontend_url(monkeypatch):
    monkeypatch.setenv("VALIDATOR_FRONTEND_URL", "https://validador.example.com")

    response = client.get("/", follow_redirects=False)

    assert response.status_code == 307
    assert response.headers["location"] == "https://validador.example.com"


def test_list_tenants_returns_display_names():
    response = client.get("/tenants", headers=auth_headers())
    assert response.status_code == 200

    payload = response.json()
    assert payload
    assert payload[0]["tenant_id"] == "default"
    assert payload[0]["display_name"] == "Default Tenant"
    assert payload[0]["is_default"] is True
    assert len(payload) == 1


def test_list_audit_events_returns_tenant_scoped_entries_in_reverse_chronological_order():
    first = audit_service.record_event(
        AuditEventType.JOB_CREATED,
        tenant_id="default",
        job_id="job-1",
        api_key_id="default",
        details={"file_name": "a.csv"},
    )
    audit_service.record_event(
        AuditEventType.JOB_COMPLETED,
        tenant_id="redesim",
        job_id="job-2",
        api_key_id="redesim",
    )
    latest = audit_service.record_event(
        AuditEventType.DUPLICATES_RESOLVED,
        tenant_id="default",
        job_id="job-3",
        api_key_id="default",
        details={"remaining_rows": 1},
    )

    response = client.get("/audit?tenant_id=default&limit=1", headers=auth_headers())
    assert response.status_code == 200

    payload = response.json()
    assert len(payload) == 1
    assert payload[0]["event_id"] == latest.event_id
    assert payload[0]["event_type"] == "duplicates_resolved"
    assert payload[0]["tenant_id"] == "default"
    assert payload[0]["job_id"] == "job-3"
    assert payload[0]["api_key_id"] == "default"
    assert payload[0]["details"]["remaining_rows"] == 1
    assert payload[0]["event_id"] != first.event_id


def test_list_audit_events_rejects_other_tenant_hint():
    response = client.get("/audit?tenant_id=redesim", headers=auth_headers())
    assert response.status_code == 403
    assert "redesim" in response.json()["detail"]


def test_platform_admin_can_read_global_audit_with_filters():
    auth_service.create_operator(
        tenant_id="default",
        username="admin.global.audit",
        password="AdminGlobal@2026",
        role=OperatorRole.PLATFORM_ADMIN,
    )
    headers, login_payload = login_headers(
        username="admin.global.audit",
        password="AdminGlobal@2026",
    )
    audit_service.record_event(
        AuditEventType.OPERATOR_CREATED,
        tenant_id="redesim",
        details={
            "actor_operator_id": login_payload["operator_id"],
            "target_operator_id": "operator-redesim",
            "result": "success",
        },
    )
    audit_service.record_event(
        AuditEventType.JOB_CREATED,
        tenant_id="default",
        job_id="job-default",
    )

    response = client.get(
        "/audit",
        headers=headers,
        params={
            "actor_operator_id": login_payload["operator_id"],
            "event_type": "operator_created",
            "created_from": (datetime.now(UTC) - timedelta(hours=1)).isoformat(),
            "created_to": (datetime.now(UTC) + timedelta(hours=1)).isoformat(),
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert len(payload) == 1
    assert payload[0]["tenant_id"] == "redesim"
    assert payload[0]["event_type"] == "operator_created"
    assert payload[0]["details"]["target_operator_id"] == "operator-redesim"


def test_tenant_admin_reads_only_own_tenant_audit_scope():
    auth_service.create_operator(
        tenant_id="default",
        username="admin.tenant.audit",
        password="AdminTenant@2026",
        role=OperatorRole.TENANT_ADMIN,
    )
    headers, login_payload = login_headers(
        username="admin.tenant.audit",
        password="AdminTenant@2026",
    )
    audit_service.record_event(
        AuditEventType.OPERATOR_CREATED,
        tenant_id="default",
        details={
            "actor_operator_id": login_payload["operator_id"],
            "target_operator_id": "operator-default",
            "result": "success",
        },
    )
    audit_service.record_event(
        AuditEventType.OPERATOR_CREATED,
        tenant_id="redesim",
        details={
            "actor_operator_id": "operator-outro-tenant",
            "target_operator_id": "operator-redesim",
            "result": "success",
        },
    )

    own_scope_response = client.get(
        "/audit",
        headers=headers,
        params={"event_type": "operator_created"},
    )
    assert own_scope_response.status_code == 200
    own_scope_payload = own_scope_response.json()
    assert own_scope_payload
    assert all(event["tenant_id"] == "default" for event in own_scope_payload)

    forbidden_response = client.get(
        "/audit",
        headers=headers,
        params={"tenant_id": "redesim"},
    )
    assert forbidden_response.status_code == 403
    assert forbidden_response.json()["detail"] == (
        "API key does not grant access to tenant 'redesim'"
    )


def test_operator_audit_view_keeps_operational_events_and_hides_admin_events():
    auth_service.create_operator(
        tenant_id="default",
        username="operador.audit",
        password="OperadorAudit@2026",
        role=OperatorRole.OPERATOR,
    )
    headers, login_payload = login_headers(
        username="operador.audit",
        password="OperadorAudit@2026",
    )
    admin_event = audit_service.record_event(
        AuditEventType.OPERATOR_CREATED,
        tenant_id="default",
        details={
            "actor_operator_id": login_payload["operator_id"],
            "target_operator_id": "operator-alvo",
            "result": "success",
        },
    )
    job_event = audit_service.record_event(
        AuditEventType.JOB_CREATED,
        tenant_id="default",
        job_id="job-operacional",
    )

    response = client.get("/audit?tenant_id=default", headers=headers)

    assert response.status_code == 200
    payload = response.json()
    returned_event_ids = {event["event_id"] for event in payload}
    assert job_event.event_id in returned_event_ids
    assert admin_event.event_id not in returned_event_ids


def test_protected_routes_require_api_key():
    response = client.get("/jobs")
    assert response.status_code == 401
    assert response.json()["detail"] == "Missing X-API-Key header"


def test_invalid_api_key_is_rejected():
    response = client.get("/jobs", headers=auth_headers("invalid-api-key"))
    assert response.status_code == 401
    assert response.json()["detail"] == "Invalid API key"


def test_initial_setup_status_is_public_and_reports_availability(tmp_path, monkeypatch):
    monkeypatch.delenv("VALIDATOR_SETUP_TOKEN", raising=False)
    enable_initial_setup_storage(monkeypatch, tmp_path)

    response = client.get("/setup")

    assert response.status_code == 200
    assert response.json() == {
        "available": True,
        "storage_configured": True,
        "requires_setup_token": False,
        "tenant_id": "default",
    }


def test_initial_setup_rejects_creation_without_persistent_storage():
    response = client.post(
        "/setup",
        json={
            "username": "admin.inicial",
            "password": "Setup@2026",
        },
    )

    assert response.status_code == 503
    assert response.json()["detail"] == "Initial setup requires persistent operator storage"


def test_initial_setup_creates_admin_once_and_then_closes_public_setup(
    tmp_path,
    monkeypatch,
):
    monkeypatch.delenv("VALIDATOR_SETUP_TOKEN", raising=False)
    operator_storage_path = enable_initial_setup_storage(monkeypatch, tmp_path)

    response = client.post(
        "/setup",
        json={
            "username": "admin.inicial",
            "password": "Setup@2026",
        },
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["tenant_id"] == "default"
    assert payload["username"] == "admin.inicial"
    assert payload["disabled"] is False
    assert payload["is_seed"] is False
    assert "password" not in payload

    assert "Setup@2026" not in operator_storage_path.read_text(encoding="utf-8")

    status_response = client.get("/setup")
    assert status_response.status_code == 200
    assert status_response.json()["available"] is False

    second_response = client.post(
        "/setup",
        json={
            "username": "outro.admin",
            "password": "Outra@2026",
        },
    )

    assert second_response.status_code == 409
    assert second_response.json()["detail"] == "Initial setup is no longer available"

    audit_events = audit_service.list_events(tenant_id="default")
    assert audit_events[0].event_type is AuditEventType.INITIAL_ADMIN_CREATED
    assert audit_events[0].details["username"] == "admin.inicial"
    assert "Setup@2026" not in json.dumps(audit_events[0].details)

    login_response = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "admin.inicial",
            "password": "Setup@2026",
        },
    )

    assert login_response.status_code == 200
    assert login_response.json()["operator_id"] == payload["operator_id"]


def test_initial_setup_accepts_required_setup_token(tmp_path, monkeypatch):
    monkeypatch.setenv("VALIDATOR_SETUP_TOKEN", "token-publicado")
    enable_initial_setup_storage(monkeypatch, tmp_path)

    status_response = client.get("/setup")
    assert status_response.status_code == 200
    assert status_response.json()["requires_setup_token"] is True

    denied_response = client.post(
        "/setup",
        json={
            "username": "admin.inicial",
            "password": "Setup@2026",
            "setup_token": "token-incorreto",
        },
    )

    assert denied_response.status_code == 403
    assert denied_response.json()["detail"] == "Invalid setup token"

    created_response = client.post(
        "/setup",
        json={
            "username": "admin.inicial",
            "password": "Setup@2026",
            "setup_token": "token-publicado",
        },
    )

    assert created_response.status_code == 201
    audit_events = audit_service.list_events(tenant_id="default")
    assert audit_events[0].event_type is AuditEventType.INITIAL_ADMIN_CREATED
    assert audit_events[0].details["setup_token_required"] is True
    assert "token-publicado" not in json.dumps(audit_events[0].details)


def test_login_emits_tenant_scoped_api_key():
    headers, payload = login_headers()

    assert payload["tenant_id"] == "default"
    assert payload["operator_id"] == "default-local-operator"
    assert payload["role"] == "tenant_admin"
    assert payload["api_key_id"].startswith("issued-")
    assert payload["x_api_key"].startswith("vapi_")
    assert payload["expires_at"]
    assert payload["header_name"] == "X-API-Key"

    stored_key = auth_service.list_records()[0]
    assert stored_key.tenant_id == "default"
    assert stored_key.operator_id == "default-local-operator"
    assert stored_key.key_hash != payload["x_api_key"]
    assert stored_key.expires_at is not None

    response = client.get("/tenants", headers=headers)
    assert response.status_code == 200
    assert response.json()[0]["tenant_id"] == "default"


def test_login_records_api_key_issue_in_audit_log():
    headers, payload = login_headers()

    response = client.get(
        "/audit",
        headers=headers,
        params={"event_type": "api_key_issued"},
    )
    assert response.status_code == 200

    audit_payload = response.json()
    assert audit_payload[0]["event_type"] == "api_key_issued"
    assert audit_payload[0]["api_key_id"] == payload["api_key_id"]
    assert audit_payload[0]["details"]["issued_ttl_seconds"] == 28800


def test_revoke_current_issued_api_key_records_audit_event_and_blocks_access():
    headers, payload = login_headers()

    response = client.post("/api-keys/revoke", headers=headers)
    assert response.status_code == 200
    revoke_payload = response.json()
    assert revoke_payload["tenant_id"] == "default"
    assert revoke_payload["api_key_id"] == payload["api_key_id"]
    assert revoke_payload["revoked_at"]

    denied = client.get("/tenants", headers=headers)
    assert denied.status_code == 401
    assert denied.json()["detail"] == "Revoked API key"

    events = audit_service.list_events(tenant_id="default")
    assert [event.event_type for event in events[:3]] == [
        AuditEventType.API_KEY_REVOKED,
        AuditEventType.ADMIN_LOGIN_SUCCEEDED,
        AuditEventType.API_KEY_ISSUED,
    ]


def test_renew_issued_api_key_rotates_session_and_records_audit_event():
    headers, payload = login_headers()

    response = client.post("/api-keys/renew", headers=headers)
    assert response.status_code == 200
    renew_payload = response.json()
    assert renew_payload["tenant_id"] == "default"
    assert renew_payload["operator_id"] == "default-local-operator"
    assert renew_payload["role"] == "tenant_admin"
    assert renew_payload["previous_api_key_id"] == payload["api_key_id"]
    assert renew_payload["api_key_id"] != payload["api_key_id"]
    assert renew_payload["x_api_key"] != payload["x_api_key"]
    assert renew_payload["expires_at"]
    assert renew_payload["header_name"] == "X-API-Key"

    denied_old = client.get("/tenants", headers=headers)
    assert denied_old.status_code == 401

    new_headers = auth_headers(renew_payload["x_api_key"])
    allowed_new = client.get("/tenants", headers=new_headers)
    assert allowed_new.status_code == 200

    events = audit_service.list_events(tenant_id="default")
    assert [event.event_type for event in events[:3]] == [
        AuditEventType.API_KEY_RENEWED,
        AuditEventType.ADMIN_LOGIN_SUCCEEDED,
        AuditEventType.API_KEY_ISSUED,
    ]
    renewal_event = events[0]
    assert renewal_event.api_key_id == payload["api_key_id"]
    assert renewal_event.details["successor_api_key_id"] == renew_payload["api_key_id"]


def test_renew_issued_api_key_rejects_expired_session():
    headers, _payload = login_headers()
    record = auth_service.list_records()[0]
    record.expires_at = datetime.now(UTC) - timedelta(seconds=1)

    response = client.post("/api-keys/renew", headers=headers)
    assert response.status_code == 401
    assert response.json()["detail"] == "Expired API key"


def test_expired_issued_api_key_returns_401_and_records_audit_event():
    headers, payload = login_headers()
    record = auth_service.list_records()[0]
    record.expires_at = datetime.now(UTC) - timedelta(seconds=1)

    response = client.get("/tenants", headers=headers)
    assert response.status_code == 401
    assert response.json()["detail"] == "Expired API key"

    events = audit_service.list_events(tenant_id="default")
    assert [event.event_type for event in events[:3]] == [
        AuditEventType.API_KEY_EXPIRED,
        AuditEventType.ADMIN_LOGIN_SUCCEEDED,
        AuditEventType.API_KEY_ISSUED,
    ]
    assert events[0].api_key_id == payload["api_key_id"]


def test_login_rejects_invalid_credentials():
    response = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "default.operator",
            "password": "wrong-password",
        },
    )
    assert response.status_code == 401
    assert response.json()["detail"] == "Invalid credentials"


def test_login_rate_limits_repeated_failures():
    for _attempt in range(4):
        response = client.post(
            "/login",
            json={
                "tenant_id": "default",
                "username": "default.operator",
                "password": "wrong-password",
            },
        )
        assert response.status_code == 401
        assert response.json()["detail"] == "Invalid credentials"

    limited = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "default.operator",
            "password": "wrong-password",
        },
    )
    assert limited.status_code == 429
    assert limited.json()["detail"] == "Too many attempts. Try again later."

    locked = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "default.operator",
            "password": DEFAULT_OPERATOR_PASSWORD,
        },
    )
    assert locked.status_code == 429
    assert locked.json()["detail"] == "Too many attempts. Try again later."


def test_login_rejects_unknown_tenant():
    response = client.post(
        "/login",
        json={
            "tenant_id": "missing",
            "username": "default.operator",
            "password": DEFAULT_OPERATOR_PASSWORD,
        },
    )
    assert response.status_code == 401
    assert response.json()["detail"] == "Invalid credentials"


def test_login_rejects_operator_for_other_tenant():
    response = client.post(
        "/login",
        json={
            "tenant_id": "redesim",
            "username": "default.operator",
            "password": DEFAULT_OPERATOR_PASSWORD,
        },
    )
    assert response.status_code == 401
    assert response.json()["detail"] == "Invalid credentials"


def test_login_accepts_redesim_v2_alias_and_returns_canonical_tenant():
    response = client.post(
        "/login",
        json={
            "tenant_id": "redesim_v2",
            "username": "redesim.operator",
            "password": "redesim-password",
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["tenant_id"] == "redesim"


def test_login_rejects_invalid_identifier_format():
    response = client.post(
        "/login",
        json={
            "tenant_id": "default';--",
            "username": "default.operator",
            "password": DEFAULT_OPERATOR_PASSWORD,
        },
    )

    assert response.status_code == 422
    assert "tenant_id contains invalid characters" in response.text


def test_issued_api_key_rejects_conflicting_tenant_hint():
    headers, _payload = login_headers(
        tenant_id="redesim",
        username="redesim.operator",
        password="redesim-password",
    )

    response = client.get("/audit?tenant_id=default", headers=headers)

    assert response.status_code == 403
    assert "default" in response.json()["detail"]


def test_list_operators_returns_tenant_scoped_summaries_without_hashes():
    auth_service.create_operator(
        tenant_id="default",
        username="novo.operador",
        password="NovaSenha@2026",
    )
    headers, _payload = login_headers()

    response = client.get(
        "/operators",
        params={"tenant_id": "default"},
        headers=headers,
    )

    assert response.status_code == 200
    payload = response.json()
    seed_operator = next(
        item for item in payload if item["operator_id"] == "default-local-operator"
    )
    created_operator = next(item for item in payload if item["username"] == "novo.operador")
    assert seed_operator["is_seed"] is True
    assert seed_operator["role"] == "tenant_admin"
    assert created_operator["disabled"] is False
    assert created_operator["role"] == "operator"
    assert created_operator["is_seed"] is False
    assert "password_hash" not in created_operator


def test_create_operator_endpoint_creates_tenant_scoped_operator():
    headers, _payload = login_headers()
    response = client.post(
        "/operators",
        headers=headers,
        json={
            "tenant_id": "default",
            "username": "novo.operador",
            "password": "NovaSenha@2026",
        },
    )

    assert response.status_code == 201
    payload = response.json()
    assert payload["tenant_id"] == "default"
    assert payload["username"] == "novo.operador"
    assert payload["role"] == "operator"
    assert payload["disabled"] is False
    assert payload["is_seed"] is False

    login_response = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "novo.operador",
            "password": "NovaSenha@2026",
        },
    )
    assert login_response.status_code == 200
    assert login_response.json()["operator_id"] == payload["operator_id"]


def test_temporary_operator_password_requires_completion_before_other_api_calls():
    headers, _payload = login_headers()
    create_response = client.post(
        "/operators",
        headers=headers,
        json={
            "tenant_id": "default",
            "username": "temporario.operador",
            "password": "Temp@2026",
        },
    )

    assert create_response.status_code == 201
    created_payload = create_response.json()
    assert created_payload["must_change_password"] is True

    login_response = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "temporario.operador",
            "password": "Temp@2026",
        },
    )
    assert login_response.status_code == 200
    login_payload = login_response.json()
    assert login_payload["must_change_password"] is True
    temp_headers = auth_headers(login_payload["x_api_key"])

    blocked = client.get("/tenants", headers=temp_headers)
    assert blocked.status_code == 403
    assert blocked.json()["detail"] == "Password change required"

    completion = client.post(
        "/operators/me/complete-password-setup",
        headers=temp_headers,
        json={"new_password": "SenhaFinal@2026"},
    )
    assert completion.status_code == 200
    assert completion.json()["must_change_password"] is False

    revoked = client.get("/tenants", headers=temp_headers)
    assert revoked.status_code == 401
    assert revoked.json()["detail"] == "Revoked API key"

    old_password = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "temporario.operador",
            "password": "Temp@2026",
        },
    )
    assert old_password.status_code == 401

    final_login = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "temporario.operador",
            "password": "SenhaFinal@2026",
        },
    )
    assert final_login.status_code == 200
    assert final_login.json()["must_change_password"] is False


def test_create_operator_invitation_accepts_once_and_does_not_persist_raw_token(
    tmp_path,
    monkeypatch,
):
    operator_storage_path = enable_initial_setup_storage(monkeypatch, tmp_path)
    admin_headers, _admin_payload = login_headers()

    create_response = client.post(
        "/operators/invitations",
        headers=admin_headers,
        json={
            "tenant_id": "default",
            "username": "convite.operador",
            "role": "operator",
            "expires_in_hours": 24,
        },
    )

    assert create_response.status_code == 201
    invitation_payload = create_response.json()
    invite_token = invitation_payload["invite_token"]
    invite_store_path = operator_storage_path.with_name("operators.invites.json")
    assert invite_store_path.exists()
    assert invite_token not in invite_store_path.read_text(encoding="utf-8")

    accept_response = client.post(
        "/operators/invitations/accept",
        json={
            "invite_token": invite_token,
            "password": "Convite@2026",
        },
    )
    assert accept_response.status_code == 201
    assert accept_response.json()["username"] == "convite.operador"

    second_accept = client.post(
        "/operators/invitations/accept",
        json={
            "invite_token": invite_token,
            "password": "OutraSenha@2026",
        },
    )
    assert second_accept.status_code == 409
    assert second_accept.json()["detail"] == "Invitation token already used"

    login_response = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "convite.operador",
            "password": "Convite@2026",
        },
    )
    assert login_response.status_code == 200
    assert login_response.json()["must_change_password"] is False


def test_password_reset_token_flow_completes_publicly_and_revokes_active_session():
    admin_headers, _admin_payload = login_headers()
    create_response = client.post(
        "/operators",
        headers=admin_headers,
        json={
            "tenant_id": "default",
            "username": "recupera.operador",
            "password": "SenhaAtual@2026",
            "require_password_change": False,
        },
    )
    assert create_response.status_code == 201
    operator_id = create_response.json()["operator_id"]

    operator_headers, _login_payload = login_headers(
        tenant_id="default",
        username="recupera.operador",
        password="SenhaAtual@2026",
    )

    token_response = client.post(
        f"/operators/{operator_id}/password-reset-token",
        headers=admin_headers,
        json={
            "tenant_id": "default",
            "expires_in_minutes": 30,
        },
    )
    assert token_response.status_code == 201
    reset_token = token_response.json()["reset_token"]

    complete_response = client.post(
        "/operators/password-reset/complete",
        json={
            "reset_token": reset_token,
            "new_password": "SenhaNova@2026",
        },
    )
    assert complete_response.status_code == 200
    assert complete_response.json()["operator_id"] == operator_id

    revoked = client.get("/tenants", headers=operator_headers)
    assert revoked.status_code == 401
    assert revoked.json()["detail"] == "Revoked API key"

    old_login = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "recupera.operador",
            "password": "SenhaAtual@2026",
        },
    )
    assert old_login.status_code == 401

    new_login = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "recupera.operador",
            "password": "SenhaNova@2026",
        },
    )
    assert new_login.status_code == 200

    reused = client.post(
        "/operators/password-reset/complete",
        json={
            "reset_token": reset_token,
            "new_password": "OutraSenha@2026",
        },
    )
    assert reused.status_code == 409
    assert reused.json()["detail"] == "Password reset token already used"


def test_disable_operator_endpoint_revokes_active_operator_session():
    admin_headers, _admin_payload = login_headers()
    create_response = client.post(
        "/operators",
        headers=admin_headers,
        json={
            "tenant_id": "default",
            "username": "ativo.operador",
            "password": "SenhaAtiva@2026",
        },
    )
    operator_id = create_response.json()["operator_id"]
    operator_headers, _login_payload = login_headers(
        tenant_id="default",
        username="ativo.operador",
        password="SenhaAtiva@2026",
    )

    response = client.post(
        f"/operators/{operator_id}/disable",
        headers=admin_headers,
        json={"tenant_id": "default"},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["operator_id"] == operator_id
    assert payload["disabled"] is True

    denied = client.get("/tenants", headers=operator_headers)
    assert denied.status_code == 401
    assert denied.json()["detail"] == "Revoked API key"


def test_reset_operator_password_endpoint_revokes_active_operator_session():
    admin_headers, _admin_payload = login_headers()
    create_response = client.post(
        "/operators",
        headers=admin_headers,
        json={
            "tenant_id": "default",
            "username": "reset.operador",
            "password": "SenhaAtual@2026",
            "require_password_change": False,
        },
    )
    assert create_response.status_code == 201
    operator_id = create_response.json()["operator_id"]

    operator_headers, _login_payload = login_headers(
        tenant_id="default",
        username="reset.operador",
        password="SenhaAtual@2026",
    )
    allowed_before_reset = client.get("/tenants", headers=operator_headers)
    assert allowed_before_reset.status_code == 200

    reset_response = client.post(
        f"/operators/{operator_id}/reset-password",
        headers=admin_headers,
        json={
            "tenant_id": "default",
            "new_password": "SenhaResetada@2026",
        },
    )
    assert reset_response.status_code == 200
    assert reset_response.json()["must_change_password"] is True

    revoked = client.get("/tenants", headers=operator_headers)
    assert revoked.status_code == 401
    assert revoked.json()["detail"] == "Revoked API key"

    old_login = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "reset.operador",
            "password": "SenhaAtual@2026",
        },
    )
    assert old_login.status_code == 401

    reset_login = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "reset.operador",
            "password": "SenhaResetada@2026",
        },
    )
    assert reset_login.status_code == 200
    assert reset_login.json()["must_change_password"] is True


def test_seed_password_rotation_endpoint_replaces_bootstrap_password():
    headers, _payload = login_headers()
    response = client.post(
        "/operators/seed/rotate-password",
        headers=headers,
        json={
            "tenant_id": "default",
            "new_password": "NovaSeed@2026",
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["operator_id"] == "default-local-operator"
    assert payload["is_seed"] is True

    old_login = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "default.operator",
            "password": DEFAULT_OPERATOR_PASSWORD,
        },
    )
    assert old_login.status_code == 401

    new_login = client.post(
        "/login",
        json={
            "tenant_id": "default",
            "username": "default.operator",
            "password": "NovaSeed@2026",
        },
    )
    assert new_login.status_code == 200

    rotated_headers = auth_headers(new_login.json()["x_api_key"])
    audit_response = client.get(
        "/audit",
        params={"tenant_id": "default"},
        headers=rotated_headers,
    )
    assert audit_response.status_code == 200
    assert any(
        event["event_type"] == "operator_password_rotated"
        for event in audit_response.json()
    )


def test_operator_management_rejects_other_tenant_scope():
    headers, _payload = login_headers()
    response = client.post(
        "/operators",
        headers=headers,
        json={
            "tenant_id": "redesim",
            "username": "bloqueado.operador",
            "password": "SenhaBloqueada@2026",
        },
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "Operator is not allowed to manage this tenant"


def test_operator_invitation_rejects_other_tenant_scope():
    headers, _payload = login_headers()
    response = client.post(
        "/operators/invitations",
        headers=headers,
        json={
            "tenant_id": "redesim",
            "username": "bloqueado.convite",
            "role": "operator",
        },
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "Operator is not allowed to manage this tenant"


def test_platform_admin_lists_tenants_and_creates_operator_in_any_tenant():
    auth_service.create_operator(
        tenant_id="default",
        username="admin.plataforma",
        password="AdminGlobal@2026",
        role=OperatorRole.PLATFORM_ADMIN,
    )
    headers, login_payload = login_headers(
        username="admin.plataforma",
        password="AdminGlobal@2026",
    )

    tenants_response = client.get("/tenants", headers=headers)
    assert tenants_response.status_code == 200
    tenant_ids = {tenant["tenant_id"] for tenant in tenants_response.json()}
    assert {"default", "redesim"}.issubset(tenant_ids)
    assert login_payload["role"] == "platform_admin"

    create_response = client.post(
        "/operators",
        headers=headers,
        json={
            "tenant_id": "redesim",
            "username": "redesim.novo",
            "password": "RedesimNovo@2026",
            "role": "operator",
        },
    )

    assert create_response.status_code == 201
    created_payload = create_response.json()
    assert created_payload["tenant_id"] == "redesim"
    assert created_payload["username"] == "redesim.novo"
    assert created_payload["role"] == "operator"

    login_response = client.post(
        "/login",
        json={
            "tenant_id": "redesim",
            "username": "redesim.novo",
            "password": "RedesimNovo@2026",
        },
    )
    assert login_response.status_code == 200
    assert login_response.json()["operator_id"] == created_payload["operator_id"]


def test_platform_admin_can_change_operator_role_and_audit_it():
    headers = platform_admin_headers()
    create_response = client.post(
        "/operators",
        headers=headers,
        json={
            "tenant_id": "redesim",
            "username": "redesim.promovido",
            "password": "Senha@2026",
            "require_password_change": False,
        },
    )
    assert create_response.status_code == 201
    operator_id = create_response.json()["operator_id"]

    update_response = client.patch(
        f"/operators/{operator_id}",
        headers=headers,
        json={
            "tenant_id": "redesim",
            "role": "tenant_admin",
        },
    )

    assert update_response.status_code == 200
    assert update_response.json()["role"] == "tenant_admin"

    event = audit_service.list_events(tenant_id="redesim")[0]
    assert event.event_type is AuditEventType.OPERATOR_ROLE_CHANGED
    assert event.details["operator_id"] == operator_id
    assert event.details["previous_role"] == "operator"
    assert event.details["new_role"] == "tenant_admin"


def test_operator_role_cannot_manage_users_even_in_own_tenant():
    auth_service.create_operator(
        tenant_id="default",
        username="operador.comum",
        password="Operador@2026",
        role=OperatorRole.OPERATOR,
    )
    headers, _payload = login_headers(
        username="operador.comum",
        password="Operador@2026",
    )

    response = client.post(
        "/operators",
        headers=headers,
        json={
            "tenant_id": "default",
            "username": "outro.operador",
            "password": "OutroOperador@2026",
        },
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "Operator is not allowed to manage this tenant"
    event = audit_service.list_events(tenant_id="default")[0]
    assert event.event_type is AuditEventType.AUTHORIZATION_DENIED
    assert event.details["operator_id"]
    assert event.details["role"] == "operator"
    assert event.details["action"] == "operators.create"


def test_tenant_admin_cannot_create_platform_admin():
    headers, _payload = login_headers()

    response = client.post(
        "/operators",
        headers=headers,
        json={
            "tenant_id": "default",
            "username": "admin.global",
            "password": "AdminGlobal@2026",
            "role": "platform_admin",
        },
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "Only platform_admin can create platform administrators"
    event = audit_service.list_events(tenant_id="default")[0]
    assert event.event_type is AuditEventType.AUTHORIZATION_DENIED
    assert event.details["operator_id"] == "default-local-operator"
    assert event.details["role"] == "tenant_admin"
    assert event.details["action"] == "operators.create.platform_admin"


def test_tenant_admin_cannot_promote_operator_to_platform_admin():
    headers, _payload = login_headers()
    create_response = client.post(
        "/operators",
        headers=headers,
        json={
            "tenant_id": "default",
            "username": "promocao.negada",
            "password": "Senha@2026",
            "require_password_change": False,
        },
    )
    assert create_response.status_code == 201
    operator_id = create_response.json()["operator_id"]

    response = client.patch(
        f"/operators/{operator_id}",
        headers=headers,
        json={
            "tenant_id": "default",
            "role": "platform_admin",
        },
    )

    assert response.status_code == 403
    assert response.json()["detail"] == "Only platform_admin can create platform administrators"
    event = audit_service.list_events(tenant_id="default")[0]
    assert event.event_type is AuditEventType.AUTHORIZATION_DENIED
    assert event.details["action"] == "operators.update_role.platform_admin"


def test_admin_tenant_routes_require_platform_admin(tmp_path, monkeypatch):
    enable_tenant_admin_storage(monkeypatch, tmp_path)
    headers, _payload = login_headers()

    response = client.get("/admin/tenants", headers=headers)

    assert response.status_code == 403
    assert response.json()["detail"] == "Only platform_admin can manage tenants"
    event = audit_service.list_events(tenant_id="default")[0]
    assert event.event_type is AuditEventType.AUTHORIZATION_DENIED
    assert event.details["action"] == "tenants.list"
    assert event.details["role"] == "tenant_admin"


def test_platform_admin_can_create_edit_disable_and_reactivate_runtime_tenant(
    tmp_path,
    monkeypatch,
):
    tenant_storage_path = enable_tenant_admin_storage(monkeypatch, tmp_path)
    headers = platform_admin_headers()

    create_response = client.post(
        "/admin/tenants",
        headers=headers,
        json={
            "tenant_id": "cliente_novo",
            "display_name": "Cliente Novo",
            "aliases": ["cliente-novo"],
        },
    )

    assert create_response.status_code == 201
    created_payload = create_response.json()
    assert created_payload["tenant_id"] == "cliente_novo"
    assert created_payload["display_name"] == "Cliente Novo"
    assert created_payload["aliases"] == ["cliente-novo"]
    assert created_payload["disabled"] is False
    assert created_payload["source"] == "runtime"

    loaded = load_tenant_config("cliente-novo")
    assert loaded.tenant_id == "cliente_novo"
    assert loaded.columns["item"] == "Item"
    assert loaded.api_keys == []
    assert loaded.operators == []

    update_response = client.patch(
        "/admin/tenants/cliente_novo",
        headers=headers,
        json={
            "display_name": "Cliente Oficial",
            "aliases": ["cliente-oficial"],
        },
    )
    assert update_response.status_code == 200
    assert update_response.json()["display_name"] == "Cliente Oficial"
    assert load_tenant_config("cliente-oficial").tenant_id == "cliente_novo"

    disable_response = client.post(
        "/admin/tenants/cliente_novo/disable",
        headers=headers,
    )
    assert disable_response.status_code == 200
    assert disable_response.json()["disabled"] is True

    hidden_from_operational_list = client.get("/tenants", headers=headers)
    assert hidden_from_operational_list.status_code == 200
    tenant_ids = {tenant["tenant_id"] for tenant in hidden_from_operational_list.json()}
    assert "cliente_novo" not in tenant_ids

    reactivate_response = client.post(
        "/admin/tenants/cliente_novo/reactivate",
        headers=headers,
    )
    assert reactivate_response.status_code == 200
    assert reactivate_response.json()["disabled"] is False
    assert load_tenant_config("cliente-oficial").tenant_id == "cliente_novo"

    admin_list_response = client.get("/admin/tenants", headers=headers)
    assert admin_list_response.status_code == 200
    assert any(
        tenant["tenant_id"] == "cliente_novo"
        and tenant["display_name"] == "Cliente Oficial"
        for tenant in admin_list_response.json()
    )

    storage_payload = json.loads(tenant_storage_path.read_text(encoding="utf-8"))
    assert storage_payload[0]["tenant_id"] == "cliente_novo"
    assert storage_payload[0]["display_name"] == "Cliente Oficial"

    event_types = [
        event.event_type
        for event in audit_service.list_events(tenant_id="cliente_novo")
        if event.event_type
        in {
            AuditEventType.TENANT_CREATED,
            AuditEventType.TENANT_UPDATED,
            AuditEventType.TENANT_DISABLED,
            AuditEventType.TENANT_REACTIVATED,
        }
    ]
    assert event_types == [
        AuditEventType.TENANT_REACTIVATED,
        AuditEventType.TENANT_DISABLED,
        AuditEventType.TENANT_UPDATED,
        AuditEventType.TENANT_CREATED,
    ]


def test_tenant_admin_create_rejects_identifier_alias_and_display_name_collisions(
    tmp_path,
    monkeypatch,
):
    enable_tenant_admin_storage(monkeypatch, tmp_path)
    headers = platform_admin_headers()

    id_collision = client.post(
        "/admin/tenants",
        headers=headers,
        json={
            "tenant_id": "default",
            "display_name": "Outro Default",
        },
    )
    assert id_collision.status_code == 409
    assert "default" in id_collision.json()["detail"]

    alias_collision = client.post(
        "/admin/tenants",
        headers=headers,
        json={
            "tenant_id": "cliente_alias",
            "display_name": "Cliente Alias",
            "aliases": ["redesim_v2"],
        },
    )
    assert alias_collision.status_code == 409
    assert "redesim_v2" in alias_collision.json()["detail"]

    display_name_collision = client.post(
        "/admin/tenants",
        headers=headers,
        json={
            "tenant_id": "cliente_display",
            "display_name": "Default Tenant",
        },
    )
    assert display_name_collision.status_code == 409
    assert display_name_collision.json()["detail"] == "Tenant display name already exists"

    invalid_identifier = client.post(
        "/admin/tenants",
        headers=headers,
        json={
            "tenant_id": "Cliente Novo",
            "display_name": "Cliente Novo",
        },
    )
    assert invalid_identifier.status_code == 422
    assert "tenant_id must use lowercase" in invalid_identifier.json()["detail"]


def test_disabled_tenant_blocks_login_and_existing_issued_keys(tmp_path, monkeypatch):
    enable_tenant_admin_storage(monkeypatch, tmp_path)
    platform_headers = platform_admin_headers()
    create_response = client.post(
        "/admin/tenants",
        headers=platform_headers,
        json={
            "tenant_id": "cliente_bloqueado",
            "display_name": "Cliente Bloqueado",
        },
    )
    assert create_response.status_code == 201
    operator = auth_service.create_operator(
        tenant_id="cliente_bloqueado",
        username="cliente.admin",
        password="Cliente@2026",
        role=OperatorRole.TENANT_ADMIN,
    )
    tenant_headers, login_payload = login_headers(
        tenant_id="cliente_bloqueado",
        username="cliente.admin",
        password="Cliente@2026",
    )
    assert login_payload["operator_id"] == operator.operator_id

    disable_response = client.post(
        "/admin/tenants/cliente_bloqueado/disable",
        headers=platform_headers,
    )
    assert disable_response.status_code == 200

    denied_existing_session = client.get("/jobs", headers=tenant_headers)
    assert denied_existing_session.status_code == 403
    assert denied_existing_session.json()["detail"] == "Tenant is disabled"

    denied_login = client.post(
        "/login",
        json={
            "tenant_id": "cliente_bloqueado",
            "username": "cliente.admin",
            "password": "Cliente@2026",
        },
    )
    assert denied_login.status_code == 401
    assert denied_login.json()["detail"] == "Invalid credentials"

    reactivate_response = client.post(
        "/admin/tenants/cliente_bloqueado/reactivate",
        headers=platform_headers,
    )
    assert reactivate_response.status_code == 200

    restored_login = client.post(
        "/login",
        json={
            "tenant_id": "cliente_bloqueado",
            "username": "cliente.admin",
            "password": "Cliente@2026",
        },
    )
    assert restored_login.status_code == 200


def test_disabled_tenant_rejects_new_operator_creation(tmp_path, monkeypatch):
    enable_tenant_admin_storage(monkeypatch, tmp_path)
    headers = platform_admin_headers()

    create_tenant = client.post(
        "/admin/tenants",
        headers=headers,
        json={
            "tenant_id": "cliente_sem_acesso",
            "display_name": "Cliente Sem Acesso",
        },
    )
    assert create_tenant.status_code == 201

    disable_tenant = client.post(
        "/admin/tenants/cliente_sem_acesso/disable",
        headers=headers,
    )
    assert disable_tenant.status_code == 200

    create_operator = client.post(
        "/operators",
        headers=headers,
        json={
            "tenant_id": "cliente_sem_acesso",
            "username": "novo.usuario",
            "password": "Senha@2026",
        },
    )
    assert create_operator.status_code == 403
    assert create_operator.json()["detail"] == "Tenant is disabled"


def test_jobs_are_scoped_to_issued_api_key_tenant():
    default_job = job_service.create_job(tenant_id="default", file_name="default.csv")
    job_service.create_job(tenant_id="redesim", file_name="redesim.csv")
    headers, _payload = login_headers()

    response = client.get("/jobs", headers=headers)

    assert response.status_code == 200
    payload = response.json()
    assert [job["job_id"] for job in payload] == [default_job.job_id]


def test_list_jobs_accepts_limit_without_leaking_other_tenant():
    older_job = job_service.create_job(tenant_id="default", file_name="older.csv")
    newer_job = job_service.create_job(tenant_id="default", file_name="newer.csv")
    job_service.create_job(tenant_id="redesim", file_name="redesim.csv")
    headers, _payload = login_headers()

    response = client.get("/jobs?limit=1", headers=headers)

    assert response.status_code == 200
    payload = response.json()
    assert [job["job_id"] for job in payload] == [newer_job.job_id]
    assert older_job.job_id not in {job["job_id"] for job in payload}


def test_upload_and_validate_creates_job():
    files = {"file": ("test.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=default",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200
    data = response.json()
    assert "job_id" in data
    assert data["tenant_id"] == "default"
    assert data["validation_scope"] == "zero_items"
    assert data["status"] == "queued"

    audit_events = [
        event for event in audit_service.list_events() if event.job_id == data["job_id"]
    ]
    assert [event.event_type.value for event in audit_events] == [
        "job_completed",
        "job_created",
    ]
    assert audit_events[0].api_key_id == DEFAULT_API_KEY_ID
    assert audit_events[1].details["file_name"] == "test.csv"
    assert audit_events[1].details["validation_scope"] == "zero_items"


def test_upload_default_tenant():
    files = {"file": ("test.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post("/validate", files=files, headers=auth_headers())
    assert response.status_code == 200
    assert response.json()["tenant_id"] == "default"
    assert response.json()["validation_scope"] == "zero_items"


def test_upload_uses_tenant_from_api_key_when_hint_missing():
    files = {
        "file": (
            "redesim.csv",
            BytesIO(REDESIM_CSV_CONTENT.encode("iso-8859-1")),
            "text/csv",
        )
    }
    response = client.post("/validate", files=files, headers=auth_headers(REDESIM_API_KEY))
    assert response.status_code == 200
    payload = response.json()
    assert payload["tenant_id"] == "redesim"

    job = job_service.get_job(payload["job_id"])
    assert job is not None
    try:
        assert job.file_path is not None
        assert job.result_path is not None
        assert job.report_path is not None
        assert Path(job.file_path).parent.name == "redesim"
        assert Path(job.result_path).parent.name == "redesim"
        assert Path(job.report_path).parent.name == "redesim"
    finally:
        if job.file_path:
            Path(job.file_path).unlink(missing_ok=True)
        if job.result_path:
            Path(job.result_path).unlink(missing_ok=True)
        if job.report_path:
            Path(job.report_path).unlink(missing_ok=True)


def test_upload_accepts_redesim_v2_alias_and_returns_canonical_tenant():
    files = {
        "file": (
            "redesim_v2.csv",
            BytesIO(REDESIM_CSV_CONTENT.encode("iso-8859-1")),
            "text/csv",
        )
    }
    response = client.post(
        "/validate?tenant_id=redesim_v2",
        files=files,
        headers=auth_headers(REDESIM_API_KEY),
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["tenant_id"] == "redesim"

    job = job_service.get_job(payload["job_id"])
    assert job is not None
    try:
        assert job.tenant_id == "redesim"
        assert job.file_path is not None
        assert Path(job.file_path).parent.name == "redesim"
    finally:
        if job.file_path:
            Path(job.file_path).unlink(missing_ok=True)
        if job.result_path:
            Path(job.result_path).unlink(missing_ok=True)
        if job.report_path:
            Path(job.report_path).unlink(missing_ok=True)


def test_upload_can_request_all_items_scope():
    files = {"file": ("test.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=default&validation_scope=all_items",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200
    assert response.json()["validation_scope"] == "all_items"


def test_upload_can_request_duplicate_items_scope():
    files = {"file": ("test.csv", BytesIO(DUPLICATE_SCOPE_CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=default&validation_scope=duplicate_items",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200
    assert response.json()["validation_scope"] == "duplicate_items"


def test_upload_can_request_llm_force_refresh():
    files = {"file": ("test.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=default&force_refresh=true",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200

    job = job_service.get_job(response.json()["job_id"])
    assert job is not None
    assert job.params[LLM_FORCE_REFRESH_PARAM] is True


def test_upload_invalid_tenant():
    files = {"file": ("test.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=nonexistent",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 403
    assert "does not grant access" in response.json()["detail"]


def test_validate_rejects_tenant_conflict_with_api_key():
    files = {"file": ("test.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=redesim",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 403
    assert "redesim" in response.json()["detail"]


def test_get_job_status():
    job = job_service.create_job(tenant_id="default")
    response = client.get(f"/jobs/{job.job_id}", headers=auth_headers())
    assert response.status_code == 200
    data = response.json()
    assert data["job_id"] == job.job_id
    assert data["status"] == "queued"
    assert data["tenant_id"] == "default"
    assert data["validation_scope"] == "zero_items"
    assert data["current_step"] == "file_received"
    assert data["status_title"] == "Arquivo recebido"
    assert data["status_detail"]
    assert data["processed_rows"] == 0
    assert data["batch_size"] == 0
    assert data["source_total_rows"] == 0
    assert data["partial_summary"] == {}
    assert data["is_partial_result_available"] is False
    assert data["partial_grouped_problems"] == {}
    assert data["partial_duplicates"] == []
    assert data["row_results_preview"] == []
    assert data["created_at"]
    assert data["updated_at"]
    assert data["cancel_requested"] is False


def test_list_jobs_can_filter_active_only():
    queued_job = job_service.create_job(tenant_id="default", file_name="queued.csv")
    running_job = job_service.create_job(tenant_id="default", file_name="running.csv")
    completed_job = job_service.create_job(tenant_id="default", file_name="done.csv")
    job_service.create_job(tenant_id="redesim", file_name="other-tenant.csv")

    running_job.mark_running()
    completed_job.mark_running()
    completed_job.mark_completed(total_rows=1)

    response = client.get("/jobs?active_only=true", headers=auth_headers())
    assert response.status_code == 200
    payload = response.json()

    assert [job["job_id"] for job in payload] == [
        running_job.job_id,
        queued_job.job_id,
    ]
    assert all(job["status"] in {"queued", "running"} for job in payload)


def test_get_job_status_rejects_other_tenant_job():
    job = job_service.create_job(tenant_id="redesim")
    response = client.get(f"/jobs/{job.job_id}", headers=auth_headers())
    assert response.status_code == 403
    assert "redesim" in response.json()["detail"]


def test_cancel_job_marks_queued_job_as_canceled():
    job = job_service.create_job(tenant_id="default", file_name="lote.csv")

    response = client.post(f"/jobs/{job.job_id}/cancel", headers=auth_headers())
    assert response.status_code == 200
    payload = response.json()

    assert payload["status"] == "canceled"
    assert payload["cancel_requested"] is False
    assert payload["status_title"] == "Processamento cancelado"


def test_cancel_job_marks_running_job_as_cancel_requested():
    job = job_service.create_job(tenant_id="default", file_name="lote.csv")
    job.mark_running()

    response = client.post(f"/jobs/{job.job_id}/cancel", headers=auth_headers())
    assert response.status_code == 200
    payload = response.json()

    assert payload["status"] == "running"
    assert payload["cancel_requested"] is True
    assert payload["status_title"] == "Cancelamento solicitado"


def test_cancel_job_reaches_canceled_after_worker_observes_request():
    job = job_service.create_job(tenant_id="default", file_name="lote.csv")
    job_service.start_job(job.job_id)

    cancel_response = client.post(
        f"/jobs/{job.job_id}/cancel",
        headers=auth_headers(),
    )
    assert cancel_response.status_code == 200
    assert cancel_response.json()["cancel_requested"] is True

    run_validation_job(job.job_id, job_service)

    status_response = client.get(f"/jobs/{job.job_id}", headers=auth_headers())
    assert status_response.status_code == 200
    payload = status_response.json()
    assert payload["status"] == "canceled"
    assert payload["cancel_requested"] is False
    assert payload["status_title"] == "Processamento cancelado"


def test_cancel_job_rejects_completed_job():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()
    job.mark_completed(total_rows=1)

    response = client.post(f"/jobs/{job.job_id}/cancel", headers=auth_headers())
    assert response.status_code == 400
    assert "queued or running" in response.json()["detail"]


def test_get_job_status_includes_file_name_metadata():
    job = job_service.create_job(tenant_id="default", file_name="lote.csv")
    response = client.get(f"/jobs/{job.job_id}", headers=auth_headers())
    assert response.status_code == 200
    data = response.json()
    assert data["file_name"] == "lote.csv"


def test_get_job_status_includes_partial_preview_payload():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()
    job_service.update_partial_result(
        job.job_id,
        total_rows=12,
        source_total_rows=20,
        processed_rows=4,
        batch_size=4,
        partial_summary={
            "total_rows": 12,
            "validated_rows": 12,
            "source_total_rows": 20,
            "processed_rows": 4,
            "rows_with_issues": 2,
            "total_issues": 3,
            "error_count": 1,
            "warning_count": 2,
        },
        partial_grouped_problems={"DUPLICATE_ITEM": [{"row_index": 0, "item": "001"}]},
        partial_duplicates=[{"item": "001", "row_indices": [0, 5], "count": 2}],
        row_results_preview=[{"row_index": 0, "item": "001", "issues": []}],
        current_step="validating_batches",
        status_title="Prévia operacional em atualização",
        status_detail="4 de 12 linhas já foram validadas.",
    )

    response = client.get(f"/jobs/{job.job_id}", headers=auth_headers())
    assert response.status_code == 200
    data = response.json()
    assert data["processed_rows"] == 4
    assert data["batch_size"] == 4
    assert data["source_total_rows"] == 20
    assert data["is_partial_result_available"] is True
    assert data["partial_summary"]["processed_rows"] == 4
    assert data["partial_summary"]["source_total_rows"] == 20
    assert data["partial_grouped_problems"]["DUPLICATE_ITEM"][0]["item"] == "001"
    assert data["partial_duplicates"][0]["count"] == 2
    assert data["row_results_preview"][0]["row_index"] == 0


def test_get_job_not_found():
    response = client.get("/jobs/nonexistent", headers=auth_headers())
    assert response.status_code == 404


def test_download_result_not_completed():
    job = job_service.create_job(tenant_id="default")
    response = client.get(f"/jobs/{job.job_id}/result", headers=auth_headers())
    assert response.status_code == 400
    assert "not completed" in response.json()["detail"]


def test_download_result_completed():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()

    with tempfile.NamedTemporaryFile(suffix=".json", delete=False, mode="w") as f:
        json.dump({"summary": {"total_rows": 1}}, f)
        result_path = f.name

    job.mark_completed(result_path=result_path, total_rows=1)

    response = client.get(f"/jobs/{job.job_id}/result", headers=auth_headers())
    assert response.status_code == 200
    assert response.json()["summary"]["total_rows"] == 1

    Path(result_path).unlink(missing_ok=True)


def test_validation_result_includes_item_and_descricao_metadata():
    files = {"file": ("test.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=default",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200

    job_id = response.json()["job_id"]
    job = job_service.get_job(job_id)
    assert job is not None
    try:
        assert job.result_path is not None
        assert job.report_path is not None
        assert job.file_path is not None

        result_response = client.get(f"/jobs/{job_id}/result", headers=auth_headers())
        assert result_response.status_code == 200

        payload = result_response.json()
        assert payload["summary"]["total_rows"] == 1
        assert payload["summary"]["source_total_rows"] == 2
        assert len(payload["row_results"]) == 1
        first_row = payload["row_results"][0]
        assert first_row["item"] == "002"
        assert first_row["descricao"] == "Cadeira"

        grouped_issue = payload["grouped_problems"]["ZERO_ITEM_COMPLEMENTO_EMPTY"][0]
        assert grouped_issue["item"] == "002"
        assert grouped_issue["descricao"] == "Cadeira"
    finally:
        if job.file_path:
            Path(job.file_path).unlink(missing_ok=True)
        if job.result_path:
            Path(job.result_path).unlink(missing_ok=True)
        if job.report_path:
            Path(job.report_path).unlink(missing_ok=True)


def test_validation_result_can_include_all_items_scope():
    files = {"file": ("test.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=default&validation_scope=all_items",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200

    job_id = response.json()["job_id"]
    job = job_service.get_job(job_id)
    assert job is not None
    try:
        result_response = client.get(f"/jobs/{job_id}/result", headers=auth_headers())
        assert result_response.status_code == 200

        payload = result_response.json()
        assert payload["summary"]["total_rows"] == 2
        assert payload["summary"]["source_total_rows"] == 2
        assert [row["item"] for row in payload["row_results"]] == ["001", "002"]
    finally:
        if job.file_path:
            Path(job.file_path).unlink(missing_ok=True)
        if job.result_path:
            Path(job.result_path).unlink(missing_ok=True)
        if job.report_path:
            Path(job.report_path).unlink(missing_ok=True)


def test_validation_result_can_include_duplicate_items_scope():
    files = {
        "file": (
            "test.csv",
            BytesIO(DUPLICATE_SCOPE_CSV_CONTENT.encode()),
            "text/csv",
        )
    }
    response = client.post(
        "/validate?tenant_id=default&validation_scope=duplicate_items",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200

    job_id = response.json()["job_id"]
    job = job_service.get_job(job_id)
    assert job is not None
    try:
        result_response = client.get(f"/jobs/{job_id}/result", headers=auth_headers())
        assert result_response.status_code == 200

        payload = result_response.json()
        assert payload["summary"]["total_rows"] == 2
        assert payload["summary"]["source_total_rows"] == 3
        assert [row["item"] for row in payload["row_results"]] == ["001", "001"]
        assert payload["duplicates"][0]["row_indices"] == [0, 2]
        assert "DUPLICATE_ITEM" in payload["grouped_problems"]
    finally:
        if job.file_path:
            Path(job.file_path).unlink(missing_ok=True)
        if job.result_path:
            Path(job.result_path).unlink(missing_ok=True)
        if job.report_path:
            Path(job.report_path).unlink(missing_ok=True)


def test_redesim_tenant_uses_configured_descricao_column_and_rules():
    files = {
        "file": (
            "redesim.csv",
            BytesIO(REDESIM_CSV_CONTENT.encode("iso-8859-1")),
            "text/csv",
        )
    }
    response = client.post(
        "/validate?tenant_id=redesim",
        files=files,
        headers=auth_headers(REDESIM_API_KEY),
    )
    assert response.status_code == 200

    job_id = response.json()["job_id"]
    job = job_service.get_job(job_id)
    assert job is not None

    try:
        result_response = client.get(
            f"/jobs/{job_id}/result",
            headers=auth_headers(REDESIM_API_KEY),
        )
        assert result_response.status_code == 200

        payload = result_response.json()
        assert payload["row_results"][0]["descricao"] == "MONITOR"
        assert "CATEGORY_MONITOR_COMPLEMENTO_REQUIRED" in payload["grouped_problems"]
        assert "CATEGORY_MONITOR_INCHES_PATTERN_MISSING" in payload["grouped_problems"]
    finally:
        if job.file_path:
            Path(job.file_path).unlink(missing_ok=True)
        if job.result_path:
            Path(job.result_path).unlink(missing_ok=True)
        if job.report_path:
            Path(job.report_path).unlink(missing_ok=True)


def test_download_report_not_completed():
    job = job_service.create_job(tenant_id="default")
    response = client.get(f"/jobs/{job.job_id}/report", headers=auth_headers())
    assert response.status_code == 400


def test_download_report_file_missing():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()
    job.mark_completed(report_path="/nonexistent/report.pdf")

    response = client.get(f"/jobs/{job.job_id}/report", headers=auth_headers())
    assert response.status_code == 404
    assert "Report file not found" in response.json()["detail"]


def test_download_report_completed():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        f.write(b"%PDF-1.4 fake content")
        report_path = f.name

    job.mark_completed(report_path=report_path)

    response = client.get(f"/jobs/{job.job_id}/report", headers=auth_headers())
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"

    Path(report_path).unlink(missing_ok=True)


def test_download_job_csv_returns_current_corrected_file():
    job = job_service.create_job(tenant_id="default", file_name="inventario.csv")

    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa corrigida,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job.file_path = csv_path

    response = client.get(f"/jobs/{job.job_id}/csv", headers=auth_headers())
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert 'filename="inventario_corrigido.csv"' in response.headers["content-disposition"]
    assert "Mesa corrigida" in response.text

    Path(csv_path).unlink(missing_ok=True)


def test_download_job_csv_file_missing():
    job = job_service.create_job(tenant_id="default", file_name="inventario.csv")
    job.file_path = "/nonexistent/inventario.csv"

    response = client.get(f"/jobs/{job.job_id}/csv", headers=auth_headers())
    assert response.status_code == 404
    assert "CSV file not found" in response.json()["detail"]


def test_download_duplicates_export_csv_for_completed_job():
    duplicate_csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
        "001,,Mesa reserva,MarcaX,ModeloY,SN2,Sala1,CC1,Detalhe,Obs\n"
    )
    files = {
        "file": (
            "duplicados.csv",
            BytesIO(duplicate_csv_content.encode()),
            "text/csv",
        )
    }
    response = client.post(
        "/validate?tenant_id=default&validation_scope=all_items",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200

    job_id = response.json()["job_id"]
    job = job_service.get_job(job_id)
    assert job is not None

    try:
        export_response = client.get(
            f"/jobs/{job_id}/exports/csv?kind=duplicates",
            headers=auth_headers(),
        )
        assert export_response.status_code == 200
        assert export_response.headers["content-type"].startswith("text/csv")
        assert 'filename="duplicados_duplicados.csv"' in (
            export_response.headers["content-disposition"]
        )
        assert "Quantidade de Ocorrências" in export_response.text
        assert "001" in export_response.text
    finally:
        if job.file_path:
            Path(job.file_path).unlink(missing_ok=True)
        if job.result_path:
            Path(job.result_path).unlink(missing_ok=True)
        if job.report_path:
            Path(job.report_path).unlink(missing_ok=True)


def test_download_problem_group_export_csv_for_completed_job():
    files = {"file": ("test.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=default",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200

    job_id = response.json()["job_id"]
    job = job_service.get_job(job_id)
    assert job is not None

    try:
        export_response = client.get(
            f"/jobs/{job_id}/exports/csv?kind=problem_group&problem_code=ZERO_ITEM_COMPLEMENTO_EMPTY",
            headers=auth_headers(),
        )
        assert export_response.status_code == 200
        assert export_response.headers["content-type"].startswith("text/csv")
        assert 'filename="test_zero_item_complemento_empty.csv"' in (
            export_response.headers["content-disposition"]
        )
        assert "Código,Linha,Item,Descrição,Severidade,Campo,Mensagem" in export_response.text
        assert "ZERO_ITEM_COMPLEMENTO_EMPTY" in export_response.text
        assert "Cadeira" in export_response.text
    finally:
        if job.file_path:
            Path(job.file_path).unlink(missing_ok=True)
        if job.result_path:
            Path(job.result_path).unlink(missing_ok=True)
        if job.report_path:
            Path(job.report_path).unlink(missing_ok=True)


def test_download_result_not_found():
    response = client.get("/jobs/nonexistent/result", headers=auth_headers())
    assert response.status_code == 404


def test_download_report_not_found():
    response = client.get("/jobs/nonexistent/report", headers=auth_headers())
    assert response.status_code == 404


def test_download_job_csv_not_found():
    response = client.get("/jobs/nonexistent/csv", headers=auth_headers())
    assert response.status_code == 404


def test_download_job_xlsx_export_for_completed_job():
    from io import BytesIO as _BytesIO

    from openpyxl import load_workbook

    xlsx_csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Cadeira,,,SN1,Sala1,CC1,,\n"
    )
    files = {"file": ("lote_excel.csv", BytesIO(xlsx_csv_content.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=default",
        files=files,
        headers=auth_headers(),
    )
    assert response.status_code == 200

    job_id = response.json()["job_id"]
    job = job_service.get_job(job_id)
    assert job is not None

    try:
        export_response = client.get(
            f"/jobs/{job_id}/export?format=xlsx",
            headers=auth_headers(),
        )
        assert export_response.status_code == 200
        assert export_response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert 'filename="lote_excel_corrigido.xlsx"' in (
            export_response.headers["content-disposition"]
        )

        workbook = load_workbook(_BytesIO(export_response.content))
        worksheet = workbook.active
        header_row = [cell.value for cell in worksheet[1]]
        assert header_row == [
            "Item",
            "Placa Anterior",
            "Descrição",
            "Marca",
            "Modelo",
            "NS",
            "Local",
            "CC",
            "Complemento",
            "Observação",
        ]
        assert worksheet.cell(row=1, column=1).font.bold is True
        assert worksheet.cell(row=2, column=1).value == "001"
    finally:
        if job.file_path:
            Path(job.file_path).unlink(missing_ok=True)
        if job.result_path:
            Path(job.result_path).unlink(missing_ok=True)
        if job.report_path:
            Path(job.report_path).unlink(missing_ok=True)


def test_download_job_xlsx_export_unsupported_format():
    job = job_service.create_job(tenant_id="default", file_name="lote.csv")
    response = client.get(
        f"/jobs/{job.job_id}/export?format=json",
        headers=auth_headers(),
    )
    assert response.status_code == 400
    assert "Unsupported export format" in response.json()["detail"]


def test_download_job_xlsx_export_not_found():
    response = client.get(
        "/jobs/nonexistent/export?format=xlsx",
        headers=auth_headers(),
    )
    assert response.status_code == 404


def test_download_job_xlsx_export_rejects_other_tenant_job():
    job = job_service.create_job(tenant_id="redesim", file_name="lote.csv")

    response = client.get(
        f"/jobs/{job.job_id}/export?format=xlsx",
        headers=auth_headers(),
    )

    assert response.status_code == 403
    assert "redesim" in response.json()["detail"]


def test_upload_saves_file():
    files = {"file": ("inventory.csv", BytesIO(CSV_CONTENT.encode()), "text/csv")}
    response = client.post(
        "/validate?tenant_id=default",
        files=files,
        headers=auth_headers(),
    )
    data = response.json()
    job = job_service.get_job(data["job_id"])
    assert job is not None
    assert job.file_path is not None
    assert job.file_name == "inventory.csv"
    assert Path(job.file_path).parent.name == "default"
    assert Path(job.file_path).name.endswith("inventory.csv")

    Path(job.file_path).unlink(missing_ok=True)
    if job.result_path:
        Path(job.result_path).unlink(missing_ok=True)
    if job.report_path:
        Path(job.report_path).unlink(missing_ok=True)


def test_job_status_shows_counters_after_completion():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()
    job.mark_completed(
        total_rows=10,
        source_total_rows=15,
        rows_with_issues=3,
        total_issues=5,
    )

    response = client.get(f"/jobs/{job.job_id}", headers=auth_headers())
    data = response.json()
    assert data["total_rows"] == 10
    assert data["source_total_rows"] == 15
    assert data["rows_with_issues"] == 3
    assert data["total_issues"] == 5
    assert data["current_step"] == "report_ready"
    assert data["status_title"] == "Relatório pronto"


def test_job_status_shows_error_after_failure():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()
    job.mark_failed("Something went wrong")

    response = client.get(f"/jobs/{job.job_id}", headers=auth_headers())
    data = response.json()
    assert data["status"] == "failed"
    assert data["error_message"] == "Something went wrong"
    assert data["current_step"] == "failed"
    assert data["status_title"] == "Falha no processamento"


def test_update_job_row_updates_csv():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()

    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job.file_path = csv_path

    response = client.patch(
        f"/jobs/{job.job_id}/rows/0",
        json={"updates": {"descricao": "Mesa executiva"}},
        headers=auth_headers(),
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["row_index"] == 0
    assert payload["updated_row"]["Descrição"] == "Mesa executiva"

    updated_csv = Path(csv_path).read_text()
    assert "Mesa executiva" in updated_csv

    Path(csv_path).unlink(missing_ok=True)


def test_get_job_row_returns_current_value_and_mapping():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()

    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job.file_path = csv_path

    response = client.get(f"/jobs/{job.job_id}/rows/0", headers=auth_headers())
    assert response.status_code == 200
    payload = response.json()
    assert payload["row"]["Descrição"] == "Mesa"
    assert payload["resolved_columns"]["descricao"] == "Descrição"

    Path(csv_path).unlink(missing_ok=True)


def test_update_job_row_review_flag_persists_and_updates_result_payload(
    tmp_path,
    monkeypatch,
):
    results_dir = tmp_path / "results"
    monkeypatch.setattr(validation_service, "RESULTS_DIR", results_dir)

    job = job_service.create_job(tenant_id="default", file_name="inventario.csv")
    result_path = tmp_path / "result.json"
    result_path.write_text(
        json.dumps(
            {
                "summary": {"total_rows": 1},
                "row_results": [
                    {
                        "row_index": 1,
                        "item": "002",
                        "descricao": "Cadeira",
                        "issues": [],
                        "has_errors": False,
                        "has_warnings": False,
                    }
                ],
                "duplicates": [],
                "grouped_problems": {
                    "ZERO_ITEM_COMPLEMENTO_EMPTY": [
                        {
                            "row_index": 1,
                            "item": "002",
                            "descricao": "Cadeira",
                            "severity": "warning",
                            "field": "complemento",
                            "message": "Complemento vazio",
                        }
                    ]
                },
            }
        ),
        encoding="utf-8",
    )
    job.mark_running()
    job.mark_completed(result_path=str(result_path), total_rows=1)

    response = client.patch(
        f"/jobs/{job.job_id}/rows/1/flag",
        json={"status": "review"},
        headers=auth_headers(),
    )
    assert response.status_code == 200
    assert response.json() == {
        "job_id": job.job_id,
        "row_index": 1,
        "status": "review",
        "review_flags": [{"row_index": 1, "status": "review"}],
    }

    flags_path = results_dir / "default" / job.job_id / "review_flags.json"
    assert flags_path.exists()
    assert json.loads(flags_path.read_text(encoding="utf-8"))["flags"] == {
        "1": "review"
    }

    result_response = client.get(f"/jobs/{job.job_id}/result", headers=auth_headers())
    assert result_response.status_code == 200
    assert result_response.json()["review_flags"] == [
        {"row_index": 1, "status": "review"}
    ]

    clear_response = client.patch(
        f"/jobs/{job.job_id}/rows/1/flag",
        json={"status": "clear"},
        headers=auth_headers(),
    )
    assert clear_response.status_code == 200
    assert clear_response.json()["review_flags"] == []


def test_update_job_row_review_flag_rejects_missing_result_row():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()

    with tempfile.NamedTemporaryFile(suffix=".json", delete=False, mode="w") as f:
        json.dump(
            {
                "summary": {"total_rows": 1},
                "row_results": [{"row_index": 0}],
                "duplicates": [],
                "grouped_problems": {},
            },
            f,
        )
        result_path = f.name

    job.mark_completed(result_path=result_path, total_rows=1)

    response = client.patch(
        f"/jobs/{job.job_id}/rows/9/flag",
        json={"status": "review"},
        headers=auth_headers(),
    )
    assert response.status_code == 400
    assert "Row index not found" in response.json()["detail"]

    Path(result_path).unlink(missing_ok=True)


def test_resolve_duplicate_rows_keeps_highest_occurrence_and_merges_missing_fields():
    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa antiga,MarcaAntiga,ModeloAntigo,SNAntigo,"
        "SalaAntiga,CCAntigo,DetalheAntigo,ObsAntiga\n"
        "001,,Mesa reserva,MarcaAnterior,ModeloAnterior,SNAnterior,"
        "SalaAnterior,CCAnterior,DetalheAnterior,ObsAnterior\n"
        "001,,Mesa atual,,ModeloAtual,,SalaAtual,CCAtual,,\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job = job_service.create_job(
        tenant_id="default",
        file_path=csv_path,
        file_name="duplicados.csv",
        params={"validation_scope": "all_items"},
    )
    run_validation_job(job.job_id, job_service)

    response = client.post(
        f"/jobs/{job.job_id}/duplicates/resolve",
        json={"row_indices": [0, 1, 2], "keep_row_index": 1},
        headers=auth_headers(),
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["kept_row_index"] == 2
    assert payload["deleted_row_indices"] == [0, 1]
    assert payload["remaining_rows"] == 1
    assert payload["merged_columns"] == [
        "Marca",
        "NS",
        "Complemento",
        "Observação",
    ]

    updated_csv = Path(csv_path).read_text(encoding="utf-8")
    assert "Mesa atual" in updated_csv
    assert "MarcaAnterior" in updated_csv
    assert "ModeloAtual" in updated_csv
    assert "SNAnterior" in updated_csv
    assert "CCAtual" in updated_csv
    assert "DetalheAnterior" in updated_csv
    assert "ObsAnterior" in updated_csv
    assert "Mesa antiga" not in updated_csv
    assert "Mesa reserva" not in updated_csv
    assert "ModeloAnterior" not in updated_csv
    assert "CCAnterior" not in updated_csv

    result_response = client.get(f"/jobs/{job.job_id}/result", headers=auth_headers())
    assert result_response.status_code == 200
    result_payload = result_response.json()
    assert result_payload["summary"]["total_rows"] == 1
    assert result_payload["duplicates"] == []
    assert [row["descricao"] for row in result_payload["row_results"]] == ["Mesa atual"]

    audit_events = [
        event for event in audit_service.list_events() if event.job_id == job.job_id
    ]
    assert [event.event_type.value for event in audit_events[:2]] == [
        "duplicates_resolved",
        "job_completed",
    ]
    assert audit_events[0].api_key_id == DEFAULT_API_KEY_ID
    assert audit_events[0].details["row_indices"] == [0, 1, 2]
    assert audit_events[0].details["kept_row_index"] == 2
    assert audit_events[0].details["deleted_row_indices"] == [0, 1]

    Path(csv_path).unlink(missing_ok=True)
    refreshed_job = job_service.get_job(job.job_id)
    if refreshed_job and refreshed_job.result_path:
        Path(refreshed_job.result_path).unlink(missing_ok=True)
    if refreshed_job and refreshed_job.report_path:
        Path(refreshed_job.report_path).unlink(missing_ok=True)


def test_resolve_duplicate_rows_same_name_skips_media_and_datetime_columns():
    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,"
        "Observação,foto_complementar_memento,data_inventario,hora_inventario,"
        "updated_at,registro,usuario\n"
        "001,,Mesa,Marca antiga,Modelo antigo,SN antigo,Sala antiga,CC antigo,"
        "Detalhe antigo,Obs antiga,https://example.com/foto-1.jpg,2026-01-01,"
        "08:00,2026-01-01T08:00:00,2026-02-01 10:15:00,Ana\n"
        "001,PA-100,Mesa,Marca intermediaria,Modelo intermediario,,"
        "Sala intermediaria,,Detalhe intermediario,,"
        "https://example.com/foto-2.jpg,2026-01-02,09:00,"
        "2026-01-02T09:00:00,2026-02-02 11:45:00,Bruno\n"
        "001,,Mesa,Marca atual,,SN atual,,CC atual,,Obs atual,,,,,,Carla\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job = job_service.create_job(
        tenant_id="default",
        file_path=csv_path,
        file_name="duplicados.csv",
        params={"validation_scope": "all_items"},
    )
    run_validation_job(job.job_id, job_service)

    response = client.post(
        f"/jobs/{job.job_id}/duplicates/resolve",
        json={"row_indices": [0, 1, 2], "keep_row_index": 2},
        headers=auth_headers(),
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["kept_row_index"] == 2
    assert payload["deleted_row_indices"] == [0, 1]
    assert payload["remaining_rows"] == 1
    assert payload["merged_columns"] == [
        "Placa Anterior",
        "Modelo",
        "Local",
        "Complemento",
    ]

    updated_csv = Path(csv_path).read_text(encoding="utf-8")
    assert "Marca atual" in updated_csv
    assert "Modelo intermediario" in updated_csv
    assert "PA-100" in updated_csv
    assert "Sala intermediaria" in updated_csv
    assert "Detalhe intermediario" in updated_csv
    assert "https://example.com/foto-1.jpg" not in updated_csv
    assert "2026-01-01T08:00:00" not in updated_csv
    assert "2026-01-02T09:00:00" not in updated_csv
    assert ",,,Carla" in updated_csv

    result_response = client.get(f"/jobs/{job.job_id}/rows/0", headers=auth_headers())
    assert result_response.status_code == 200
    row_payload = result_response.json()["row"]
    assert row_payload["foto_complementar_memento"] == ""
    assert row_payload["data_inventario"] == ""
    assert row_payload["hora_inventario"] == ""
    assert row_payload["updated_at"] == ""
    assert row_payload["registro"] == ""
    assert row_payload["usuario"] == "Carla"

    Path(csv_path).unlink(missing_ok=True)
    refreshed_job = job_service.get_job(job.job_id)
    if refreshed_job and refreshed_job.result_path:
        Path(refreshed_job.result_path).unlink(missing_ok=True)
    if refreshed_job and refreshed_job.report_path:
        Path(refreshed_job.report_path).unlink(missing_ok=True)


def test_resolve_duplicate_rows_requires_completed_job_for_in_place_refresh():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()

    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
        "001,,Mesa reserva,MarcaX,ModeloY,SN2,Sala1,CC1,Detalhe,Obs\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job.file_path = csv_path

    response = client.post(
        f"/jobs/{job.job_id}/duplicates/resolve",
        json={"row_indices": [0, 1], "keep_row_index": 1},
        headers=auth_headers(),
    )
    assert response.status_code == 400
    assert "completed jobs" in response.json()["detail"]

    unchanged_csv = Path(csv_path).read_text(encoding="utf-8")
    assert "Mesa reserva" in unchanged_csv
    assert "Mesa," in unchanged_csv

    Path(csv_path).unlink(missing_ok=True)


def test_resolve_duplicate_rows_rejects_keep_index_outside_group():
    job = job_service.create_job(tenant_id="default")
    job.mark_running()

    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
        "001,,Mesa reserva,MarcaX,ModeloY,SN2,Sala1,CC1,Detalhe,Obs\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job.file_path = csv_path

    response = client.post(
        f"/jobs/{job.job_id}/duplicates/resolve",
        json={"row_indices": [0, 1], "keep_row_index": 3},
        headers=auth_headers(),
    )
    assert response.status_code == 400
    assert "keep_row_index" in response.json()["detail"]

    Path(csv_path).unlink(missing_ok=True)


def test_reprocess_job_creates_new_job_from_corrected_csv():
    job = job_service.create_job(tenant_id="default", file_name="corrigido.csv")
    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa executiva,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job.file_path = csv_path

    response = client.post(f"/jobs/{job.job_id}/reprocess", headers=auth_headers())
    assert response.status_code == 200
    payload = response.json()
    assert payload["job_id"] != job.job_id
    assert payload["tenant_id"] == "default"
    assert payload["validation_scope"] == "zero_items"

    new_job = job_service.get_job(payload["job_id"])
    assert new_job is not None
    assert new_job.file_path is not None
    assert new_job.file_path != csv_path
    assert Path(new_job.file_path).parent.name == "default"
    assert "Mesa executiva" in Path(new_job.file_path).read_text()

    Path(csv_path).unlink(missing_ok=True)
    Path(new_job.file_path).unlink(missing_ok=True)
    if new_job.result_path:
        Path(new_job.result_path).unlink(missing_ok=True)
    if new_job.report_path:
        Path(new_job.report_path).unlink(missing_ok=True)


def test_reprocess_job_preserves_validation_scope():
    job = job_service.create_job(
        tenant_id="default",
        file_name="corrigido.csv",
        params={"validation_scope": "all_items"},
    )
    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa executiva,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job.file_path = csv_path

    response = client.post(f"/jobs/{job.job_id}/reprocess", headers=auth_headers())
    assert response.status_code == 200
    payload = response.json()
    assert payload["validation_scope"] == "all_items"

    new_job = job_service.get_job(payload["job_id"])
    assert new_job is not None
    assert new_job.params["validation_scope"] == "all_items"

    Path(csv_path).unlink(missing_ok=True)
    if new_job.file_path:
        Path(new_job.file_path).unlink(missing_ok=True)
    if new_job.result_path:
        Path(new_job.result_path).unlink(missing_ok=True)
    if new_job.report_path:
        Path(new_job.report_path).unlink(missing_ok=True)


def test_reprocess_job_can_override_llm_force_refresh():
    job = job_service.create_job(
        tenant_id="default",
        file_name="corrigido.csv",
        params={"validation_scope": "all_items", LLM_FORCE_REFRESH_PARAM: False},
    )
    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa executiva,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job.file_path = csv_path

    response = client.post(
        f"/jobs/{job.job_id}/reprocess?force_refresh=true",
        headers=auth_headers(),
    )
    assert response.status_code == 200

    new_job = job_service.get_job(response.json()["job_id"])
    assert new_job is not None
    assert new_job.params["validation_scope"] == "all_items"
    assert new_job.params[LLM_FORCE_REFRESH_PARAM] is True

    Path(csv_path).unlink(missing_ok=True)
    if new_job.file_path:
        Path(new_job.file_path).unlink(missing_ok=True)
    if new_job.result_path:
        Path(new_job.result_path).unlink(missing_ok=True)
    if new_job.report_path:
        Path(new_job.report_path).unlink(missing_ok=True)


def test_reprocess_job_persists_and_exposes_lineage_and_audit():
    job = job_service.create_job(tenant_id="default", file_name="corrigido.csv")
    csv_content = (
        "Item,Placa Anterior,Descrição,Marca,Modelo,NS,Local,CC,Complemento,Observação\n"
        "001,,Mesa executiva,MarcaX,ModeloY,SN1,Sala1,CC1,Detalhe,Obs\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    job.file_path = csv_path

    response = client.post(f"/jobs/{job.job_id}/reprocess", headers=auth_headers())
    assert response.status_code == 200
    new_job_id = response.json()["job_id"]

    new_status = client.get(f"/jobs/{new_job_id}", headers=auth_headers())
    assert new_status.status_code == 200
    assert new_status.json()["parent_job_id"] == job.job_id
    assert new_status.json()["latest_retry_job_id"] is None

    source_status = client.get(f"/jobs/{job.job_id}", headers=auth_headers())
    assert source_status.status_code == 200
    assert source_status.json()["parent_job_id"] is None
    assert source_status.json()["latest_retry_job_id"] == new_job_id

    reprocess_events = [
        event
        for event in audit_service.list_events()
        if event.event_type == AuditEventType.JOB_REPROCESSED
    ]
    assert len(reprocess_events) == 1
    assert reprocess_events[0].job_id == job.job_id
    assert reprocess_events[0].details["new_job_id"] == new_job_id
    assert reprocess_events[0].details["parent_job_id"] == job.job_id

    new_job_created = [
        event
        for event in audit_service.list_events()
        if event.event_type == AuditEventType.JOB_CREATED and event.job_id == new_job_id
    ]
    assert new_job_created
    assert new_job_created[0].details["parent_job_id"] == job.job_id

    new_job = job_service.get_job(new_job_id)
    Path(csv_path).unlink(missing_ok=True)
    if new_job and new_job.file_path:
        Path(new_job.file_path).unlink(missing_ok=True)
    if new_job and new_job.result_path:
        Path(new_job.result_path).unlink(missing_ok=True)
    if new_job and new_job.report_path:
        Path(new_job.report_path).unlink(missing_ok=True)


def test_reprocess_lineage_respects_tenant_authorization():
    other_tenant_job = job_service.create_job(
        tenant_id="redesim",
        file_name="outro.csv",
    )
    csv_content = (
        "Item,Descrição,Marca,Modelo,Complemento,NS\n"
        "001,MONITOR,Dell,P2419H,,SN1\n"
    )
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w") as f:
        f.write(csv_content)
        csv_path = f.name

    other_tenant_job.file_path = csv_path

    response = client.get(
        f"/jobs/{other_tenant_job.job_id}",
        headers=auth_headers(),
    )
    assert response.status_code == 403

    response = client.post(
        f"/jobs/{other_tenant_job.job_id}/reprocess",
        headers=auth_headers(),
    )
    assert response.status_code == 403

    Path(csv_path).unlink(missing_ok=True)

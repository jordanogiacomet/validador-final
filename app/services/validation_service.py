import csv
import json
import math
import os
import re
import shutil
import time
import unicodedata
from collections.abc import Callable
from concurrent.futures import FIRST_COMPLETED, Future, ThreadPoolExecutor, wait
from dataclasses import dataclass
from datetime import UTC, datetime
from enum import StrEnum
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.core.canonical_fields import (
    DEFAULT_TENANT_COLUMNS,
    resolve_source_column_name,
)
from app.core.engine import ValidationEngine, ValidationExecutionPlan
from app.core.issue import ValidationIssue
from app.core.job import JobRecord, JobStatus
from app.core.llm_cache import LLM_FORCE_REFRESH_PARAM, resolve_force_refresh
from app.core.logging import get_logger, log_event
from app.core.registry import register_rule
from app.core.tenant_config import TenantConfig
from app.core.tenant_loader import canonicalize_tenant_id, load_tenant_config
from app.core.validation_scope import (
    VALIDATION_SCOPE_PARAM,
    ValidationScope,
    parse_validation_scope,
)
from app.rules.brand_model_consistency import BrandModelConsistencyRule
from app.rules.category_rules import CategoryCriticalCheckRule, CategoryRequiredFieldsRule
from app.rules.integrity import DuplicateItemRule, FlagConsistencyRule
from app.rules.llm_audit import (
    LLM_AUDIT_METADATA_KEY,
    LLMAuditRule,
    merge_llm_audit_metadata,
    snapshot_llm_audit_metadata,
)
from app.rules.suspicious_patterns import SuspiciousPatternRule
from app.rules.zero_item_quality import ZeroItemQualityRule
from app.services.job_service import DEFAULT_EXECUTION_LEASE_SECONDS, JobService
from app.services.report_service import (
    build_duplicate_section,
    build_duplicates_export_csv,
    build_full_report,
    build_partial_report,
    build_problem_group_export_csv,
    build_xlsx_export_bytes,
    generate_pdf_report,
)

UPLOADS_DIR = Path("uploads")
RESULTS_DIR = Path("results")
DEFAULT_BATCH_SIZE = 200
TARGET_PREVIEW_BATCHES = 10
REVIEW_FLAGS_FILE_NAME = "review_flags.json"
CSV_READ_CHUNK_SIZE = 5000
COOPERATIVE_CHECKPOINT_ROW_INTERVAL = 1000
PARALLEL_CANCELLATION_POLL_SECONDS = 0.1
UPLOAD_MAX_BYTES_ENV = "VALIDATOR_UPLOAD_MAX_BYTES"
DEFAULT_UPLOAD_MAX_BYTES = 10 * 1024 * 1024
SUPPORTED_UPLOAD_SUFFIXES = {".csv", ".xlsx"}
COMMON_CSV_DELIMITERS = (",", ";", "\t", "|")
COMMON_CSV_ENCODINGS = ("utf-8", "utf-8-sig", "iso-8859-1")
UPLOAD_TEMPLATE_DATA_SHEET_NAME = "Planilha"
UPLOAD_TEMPLATE_GUIDE_SHEET_NAME = "Instrucoes"

_logger = get_logger("validation_service")
CancellationCheckpoint = Callable[[], None]


@dataclass
class _CachedJobCsvContext:
    tenant_id: str
    file_path: Path
    signature: tuple[int, int]
    tenant_config: TenantConfig
    df: pd.DataFrame


@dataclass(frozen=True)
class UploadPreflightIssue:
    code: str
    message: str


@dataclass(frozen=True)
class UploadPreflightResult:
    file_name: str
    file_size_bytes: int
    detected_columns: list[str]
    missing_columns: list[str]
    guidance: list[str]
    issues: list[UploadPreflightIssue]


@dataclass(frozen=True)
class _CsvLayoutProbe:
    encoding: str
    delimiter: str
    columns: list[str]
    missing_columns: list[str]
    matched_column_count: int
    decode_error: str | None = None


_JOB_CSV_CONTEXT_CACHE: dict[str, _CachedJobCsvContext] = {}

MEDIA_COLUMN_NAMES = {
    "foto",
    "foto_complementar",
    "foto_complementar_memento",
    "imagem",
    "image",
    "photo",
    "media",
    "midia",
}
MEDIA_COLUMN_TOKENS = {
    "foto",
    "imagem",
    "image",
    "photo",
    "media",
    "midia",
    "memento",
}
DATE_TIME_COLUMN_TOKENS = {
    "data",
    "date",
    "hora",
    "hour",
    "time",
    "timestamp",
    "datetime",
    "dt",
    "dh",
}
DATE_TIME_VALUE_PATTERNS = (
    re.compile(
        r"^\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}(?::\d{2})?(?:\.\d+)?)?"
        r"(?:Z|[+-]\d{2}:\d{2})?$"
    ),
    re.compile(r"^\d{2}/\d{2}/\d{4}(?: \d{2}:\d{2}(?::\d{2})?)?$"),
    re.compile(r"^\d{2}-\d{2}-\d{4}(?: \d{2}:\d{2}(?::\d{2})?)?$"),
    re.compile(r"^\d{1,2}:\d{2}(?::\d{2})?$"),
)


def _sanitize_job_file_name(
    file_name: str | None,
    *,
    default: str = "lote.csv",
) -> str:
    normalized_name = Path(file_name or default).name
    return normalized_name or default


def _get_upload_extension(file_name: str | None) -> str:
    return Path(_sanitize_job_file_name(file_name)).suffix.lower()


def _build_storage_csv_file_name(file_name: str | None) -> str:
    sanitized_name = _sanitize_job_file_name(file_name)
    stem = Path(sanitized_name).stem or "lote"
    return f"{stem}.csv"


def get_tenant_uploads_dir(tenant_id: str) -> Path:
    return UPLOADS_DIR / canonicalize_tenant_id(tenant_id)


def get_tenant_results_dir(tenant_id: str) -> Path:
    return RESULTS_DIR / canonicalize_tenant_id(tenant_id)


def build_job_upload_path(
    tenant_id: str,
    job_id: str,
    file_name: str | None,
) -> Path:
    normalized_name = _sanitize_job_file_name(file_name)
    return get_tenant_uploads_dir(tenant_id) / f"{job_id}_{normalized_name}"


def build_job_result_path(tenant_id: str, job_id: str) -> Path:
    return get_tenant_results_dir(tenant_id) / f"{job_id}_result.json"


def build_job_report_path(tenant_id: str, job_id: str) -> Path:
    return get_tenant_results_dir(tenant_id) / f"{job_id}_report.pdf"


def build_job_review_flags_path(tenant_id: str, job_id: str) -> Path:
    return get_tenant_results_dir(tenant_id) / job_id / REVIEW_FLAGS_FILE_NAME


def _build_legacy_upload_path(
    job_id: str,
    file_name: str | None,
) -> Path:
    normalized_name = _sanitize_job_file_name(file_name)
    return UPLOADS_DIR / f"{job_id}_{normalized_name}"


def _build_legacy_result_path(job_id: str) -> Path:
    return RESULTS_DIR / f"{job_id}_result.json"


def _build_legacy_report_path(job_id: str) -> Path:
    return RESULTS_DIR / f"{job_id}_report.pdf"


def _resolve_existing_artifact_path(
    *,
    persisted_path: str | None,
    preferred_path: Path,
    legacy_path: Path | None = None,
) -> Path:
    candidates: list[Path] = []
    if persisted_path:
        candidates.append(Path(persisted_path))

    if preferred_path not in candidates:
        candidates.append(preferred_path)

    if legacy_path is not None and legacy_path not in candidates:
        candidates.append(legacy_path)

    for candidate in candidates:
        if candidate.exists():
            return candidate

    if persisted_path:
        return Path(persisted_path)

    return preferred_path


def _resolve_job_upload_path(job: JobRecord) -> Path:
    return _resolve_existing_artifact_path(
        persisted_path=job.file_path,
        preferred_path=build_job_upload_path(job.tenant_id, job.job_id, job.file_name),
        legacy_path=_build_legacy_upload_path(job.job_id, job.file_name),
    )


def _resolve_job_result_path(job: JobRecord) -> Path:
    return _resolve_existing_artifact_path(
        persisted_path=job.result_path,
        preferred_path=build_job_result_path(job.tenant_id, job.job_id),
        legacy_path=_build_legacy_result_path(job.job_id),
    )


def _resolve_job_report_path(job: JobRecord) -> Path:
    return _resolve_existing_artifact_path(
        persisted_path=job.report_path,
        preferred_path=build_job_report_path(job.tenant_id, job.job_id),
        legacy_path=_build_legacy_report_path(job.job_id),
    )


def _resolve_job_review_flags_path(job: JobRecord) -> Path:
    return build_job_review_flags_path(job.tenant_id, job.job_id)


def resolve_job_upload_path(job: JobRecord) -> Path:
    return _resolve_job_upload_path(job)


def resolve_job_result_path(job: JobRecord) -> Path:
    return _resolve_job_result_path(job)


def resolve_job_report_path(job: JobRecord) -> Path:
    return _resolve_job_report_path(job)


def resolve_job_review_flags_path(job: JobRecord) -> Path:
    return _resolve_job_review_flags_path(job)


def _build_csv_read_kwargs(tenant_config: TenantConfig) -> dict[str, str | bool]:
    return {
        "sep": tenant_config.csv.delimiter,
        "encoding": tenant_config.csv.encoding,
        "keep_default_na": False,
    }


def _display_csv_delimiter(delimiter: str) -> str:
    if delimiter == "\t":
        return "TAB"
    return delimiter


def _format_upload_size_limit(max_bytes: int) -> str:
    if max_bytes >= 1024 * 1024:
        size_mb = max_bytes / (1024 * 1024)
        return f"{int(size_mb)} MB" if size_mb.is_integer() else f"{size_mb:.1f} MB"

    if max_bytes >= 1024:
        size_kb = max_bytes / 1024
        return f"{int(size_kb)} KB" if size_kb.is_integer() else f"{size_kb:.1f} KB"

    return f"{max_bytes} bytes"


def _get_upload_max_bytes() -> int:
    configured_value = os.getenv(UPLOAD_MAX_BYTES_ENV, "").strip()
    if not configured_value:
        return DEFAULT_UPLOAD_MAX_BYTES

    try:
        parsed_value = int(configured_value)
    except ValueError:
        return DEFAULT_UPLOAD_MAX_BYTES

    return max(parsed_value, 1)


def _sanitize_header_name(value: object) -> str:
    return str(value).removeprefix("\ufeff").strip()


def _trim_trailing_empty_cells(values: list[str]) -> list[str]:
    trimmed_values = list(values)
    while trimmed_values and not trimmed_values[-1].strip():
        trimmed_values.pop()
    return trimmed_values


def _stringify_spreadsheet_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_xlsx_content(content: bytes) -> tuple[list[str], pd.DataFrame]:
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:  # pragma: no cover - defensive branch
        raise ValueError("Unable to read XLSX content.") from exc

    try:
        worksheets = workbook.worksheets
        if not worksheets:
            return [], pd.DataFrame()

        worksheet = worksheets[0]
        raw_headers: list[str] | None = None
        data_rows: list[list[str]] = []
        header_width = 0

        for raw_row in worksheet.iter_rows(values_only=True):
            row_values = _trim_trailing_empty_cells(
                [_stringify_spreadsheet_cell(value) for value in raw_row]
            )
            if raw_headers is None:
                sanitized_headers = [_sanitize_header_name(value) for value in row_values]
                if not any(sanitized_headers):
                    continue
                raw_headers = sanitized_headers
                header_width = len(raw_headers)
                continue

            if not any(cell.strip() for cell in row_values):
                continue

            normalized_row = row_values[:header_width]
            if len(normalized_row) < header_width:
                normalized_row.extend([""] * (header_width - len(normalized_row)))
            data_rows.append(normalized_row)

        if raw_headers is None:
            return [], pd.DataFrame()

        return raw_headers, pd.DataFrame(data_rows, columns=raw_headers)
    finally:
        workbook.close()


def _get_expected_upload_columns(tenant_config: TenantConfig) -> list[str]:
    expected_columns: list[str] = []
    seen_columns: set[str] = set()

    for field_name, default_column in DEFAULT_TENANT_COLUMNS.items():
        source_column = str(tenant_config.columns.get(field_name, default_column)).strip()
        if not source_column or source_column in seen_columns:
            continue
        expected_columns.append(source_column)
        seen_columns.add(source_column)

    return expected_columns


def _probe_csv_layout(
    content: bytes,
    *,
    tenant_config: TenantConfig,
    encoding: str,
    delimiter: str,
) -> _CsvLayoutProbe:
    expected_columns = _get_expected_upload_columns(tenant_config)

    try:
        decoded_text = content.decode(encoding)
    except UnicodeDecodeError as exc:
        return _CsvLayoutProbe(
            encoding=encoding,
            delimiter=delimiter,
            columns=[],
            missing_columns=expected_columns,
            matched_column_count=0,
            decode_error=str(exc),
        )

    reader = csv.reader(StringIO(decoded_text), delimiter=delimiter)
    header_row: list[str] = []
    for row in reader:
        cleaned_row = [_sanitize_header_name(cell) for cell in row]
        if any(cleaned_row):
            header_row = [cell for cell in cleaned_row if cell]
            break

    missing_columns = [
        source_column
        for source_column in expected_columns
        if resolve_source_column_name(source_column, header_row) is None
    ]

    return _CsvLayoutProbe(
        encoding=encoding,
        delimiter=delimiter,
        columns=header_row,
        missing_columns=missing_columns,
        matched_column_count=len(expected_columns) - len(missing_columns),
    )


def _match_expected_upload_columns(
    expected_columns: list[str],
    detected_columns: list[str],
) -> tuple[list[str], int]:
    missing_columns: list[str] = []
    matched_column_count = 0

    for source_column in expected_columns:
        if resolve_source_column_name(source_column, detected_columns) is None:
            missing_columns.append(source_column)
            continue
        matched_column_count += 1

    return missing_columns, matched_column_count


def _find_best_alternative_csv_probe(
    content: bytes,
    *,
    tenant_config: TenantConfig,
    expected_encoding: str,
    expected_delimiter: str,
) -> _CsvLayoutProbe | None:
    candidate_probes: list[_CsvLayoutProbe] = []
    seen_combinations: set[tuple[str, str]] = set()

    candidate_encodings = [expected_encoding]
    candidate_encodings.extend(
        encoding
        for encoding in COMMON_CSV_ENCODINGS
        if encoding not in candidate_encodings
    )

    candidate_delimiters = [expected_delimiter]
    candidate_delimiters.extend(
        delimiter
        for delimiter in COMMON_CSV_DELIMITERS
        if delimiter not in candidate_delimiters
    )

    for encoding in candidate_encodings:
        for delimiter in candidate_delimiters:
            if encoding == expected_encoding and delimiter == expected_delimiter:
                continue

            combination = (encoding, delimiter)
            if combination in seen_combinations:
                continue
            seen_combinations.add(combination)

            probe = _probe_csv_layout(
                content,
                tenant_config=tenant_config,
                encoding=encoding,
                delimiter=delimiter,
            )
            if probe.decode_error is None:
                candidate_probes.append(probe)

    if not candidate_probes:
        return None

    return max(
        candidate_probes,
        key=lambda probe: (probe.matched_column_count, len(probe.columns)),
    )


def _build_preflight_result(
    *,
    file_name: str,
    file_size_bytes: int,
    detected_columns: list[str],
    missing_columns: list[str],
    guidance: list[str],
    issues: list[UploadPreflightIssue],
) -> UploadPreflightResult:
    return UploadPreflightResult(
        file_name=file_name,
        file_size_bytes=file_size_bytes,
        detected_columns=detected_columns,
        missing_columns=missing_columns,
        guidance=guidance,
        issues=issues,
    )


def _run_xlsx_upload_preflight(
    *,
    file_name: str,
    file_size_bytes: int,
    content: bytes,
    tenant_config: TenantConfig,
) -> UploadPreflightResult:
    expected_columns = _get_expected_upload_columns(tenant_config)

    try:
        raw_headers, df = _read_xlsx_content(content)
    except ValueError:
        _raise_upload_preflight_error(
            "Nao foi possivel ler a planilha XLSX enviada.",
            file_name=file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=[],
            missing_columns=expected_columns,
            guidance=[
                "Baixe novamente o modelo da empresa ou reexporte a planilha em XLSX.",
                "Se preferir, envie o arquivo como CSV com o layout configurado para a empresa.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="xlsx_unreadable",
                    message="A leitura da planilha XLSX falhou antes da criacao do job.",
                )
            ],
        )

    detected_columns = [column for column in raw_headers if column]
    missing_columns, matched_column_count = _match_expected_upload_columns(
        expected_columns,
        detected_columns,
    )

    if not detected_columns and df.empty:
        _raise_upload_preflight_error(
            "O arquivo enviado esta vazio.",
            file_name=file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=[],
            missing_columns=expected_columns,
            guidance=[
                "Confirme se a planilha tem cabecalho na primeira linha.",
                "Inclua ao menos uma linha de dados antes de reenviar o arquivo.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="file_empty",
                    message="Nenhum conteudo util foi encontrado na planilha XLSX.",
                )
            ],
        )

    if not detected_columns or matched_column_count == 0:
        _raise_upload_preflight_error(
            "A primeira linha nao contem um cabecalho compativel com o layout esperado.",
            file_name=file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=detected_columns,
            missing_columns=expected_columns,
            guidance=[
                "Confirme se a primeira linha traz o cabecalho da planilha.",
                "Use o modelo da empresa para manter os nomes das colunas esperadas.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="missing_header",
                    message="Nenhuma coluna obrigatoria foi reconhecida na primeira linha.",
                )
            ],
        )

    if missing_columns:
        _raise_upload_preflight_error(
            "Faltam colunas obrigatorias no cabecalho da planilha.",
            file_name=file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=detected_columns,
            missing_columns=missing_columns,
            guidance=[
                "Inclua as colunas faltantes no cabecalho antes de reenviar.",
                "Use o modelo da empresa para manter a ordem e os nomes esperados.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="missing_columns",
                    message=(
                        "O cabecalho foi lido, mas ainda nao contem todas "
                        "as colunas necessarias."
                    ),
                )
            ],
        )

    if df.empty:
        _raise_upload_preflight_error(
            "O arquivo nao contem linhas de dados para validar.",
            file_name=file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=detected_columns,
            missing_columns=[],
            guidance=[
                "Mantenha o cabecalho na primeira linha e preencha ao menos um item.",
                "Reexporte a planilha depois de conferir que ha registros abaixo do cabecalho.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="missing_rows",
                    message="Nenhuma linha de dados foi encontrada abaixo do cabecalho.",
                )
            ],
        )

    return _build_preflight_result(
        file_name=file_name,
        file_size_bytes=file_size_bytes,
        detected_columns=detected_columns,
        missing_columns=[],
        guidance=[],
        issues=[],
    )


def _raise_upload_preflight_error(
    detail: str,
    *,
    file_name: str,
    file_size_bytes: int,
    detected_columns: list[str],
    missing_columns: list[str],
    guidance: list[str],
    issues: list[UploadPreflightIssue],
) -> None:
    raise UploadPreflightError(
        detail,
        _build_preflight_result(
            file_name=file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=detected_columns,
            missing_columns=missing_columns,
            guidance=guidance,
            issues=issues,
        ),
    )


def run_upload_preflight(
    *,
    file_name: str | None,
    content: bytes,
    tenant_config: TenantConfig,
) -> UploadPreflightResult:
    sanitized_file_name = _sanitize_job_file_name(file_name)
    file_size_bytes = len(content)
    file_extension = _get_upload_extension(sanitized_file_name)
    expected_columns = _get_expected_upload_columns(tenant_config)

    if file_extension not in SUPPORTED_UPLOAD_SUFFIXES:
        _raise_upload_preflight_error(
            "Envie a planilha em CSV ou XLSX antes de iniciar o lote.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=[],
            missing_columns=expected_columns,
            guidance=[
                "Exporte o arquivo novamente em CSV ou XLSX.",
                "Outros formatos ainda nao sao aceitos nesta etapa.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="unsupported_extension",
                    message=f"Formato '{file_extension or 'sem extensao'}' nao suportado.",
                )
            ],
        )

    if not content or not content.strip():
        _raise_upload_preflight_error(
            "O arquivo enviado esta vazio.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=[],
            missing_columns=expected_columns,
            guidance=[
                "Confirme se a planilha tem cabecalho na primeira linha.",
                "Inclua ao menos uma linha de dados antes de reenviar o arquivo.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="file_empty",
                    message="Nenhum conteudo util foi encontrado no arquivo.",
                )
            ],
        )

    max_upload_bytes = _get_upload_max_bytes()
    if file_size_bytes > max_upload_bytes:
        _raise_upload_preflight_error(
            "O arquivo excede o tamanho maximo permitido para preflight.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=[],
            missing_columns=expected_columns,
            guidance=[
                f"Reduza o arquivo para ate {_format_upload_size_limit(max_upload_bytes)}.",
                "Se necessario, divida o lote em partes menores antes do upload.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="file_too_large",
                    message=(
                        f"O arquivo tem {file_size_bytes} bytes e o limite atual eh "
                        f"{max_upload_bytes} bytes."
                    ),
                )
            ],
        )

    if file_extension == ".xlsx":
        return _run_xlsx_upload_preflight(
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            content=content,
            tenant_config=tenant_config,
        )

    expected_probe = _probe_csv_layout(
        content,
        tenant_config=tenant_config,
        encoding=tenant_config.csv.encoding,
        delimiter=tenant_config.csv.delimiter,
    )
    best_alternative_probe = _find_best_alternative_csv_probe(
        content,
        tenant_config=tenant_config,
        expected_encoding=tenant_config.csv.encoding,
        expected_delimiter=tenant_config.csv.delimiter,
    )

    if expected_probe.decode_error is not None:
        guidance = [
            f"Salve o CSV com codificacao {tenant_config.csv.encoding}.",
            "Reexporte o arquivo antes de tentar novamente.",
        ]
        if (
            best_alternative_probe is not None
            and best_alternative_probe.encoding != tenant_config.csv.encoding
            and best_alternative_probe.matched_column_count > 0
        ):
            guidance.append(
                f"O arquivo atual parece estar em {best_alternative_probe.encoding}."
            )
        _raise_upload_preflight_error(
            "A codificacao do arquivo nao corresponde ao layout esperado para esta empresa.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=[],
            missing_columns=expected_columns,
            guidance=guidance,
            issues=[
                UploadPreflightIssue(
                    code="encoding_mismatch",
                    message=(
                        f"Falha ao ler o arquivo com codificacao "
                        f"{tenant_config.csv.encoding}."
                    ),
                )
            ],
        )

    if (
        best_alternative_probe is not None
        and best_alternative_probe.matched_column_count
        > expected_probe.matched_column_count
    ):
        detected_columns = best_alternative_probe.columns

        if (
            best_alternative_probe.delimiter != tenant_config.csv.delimiter
            and best_alternative_probe.encoding == tenant_config.csv.encoding
        ):
            _raise_upload_preflight_error(
                "O delimitador do CSV nao corresponde ao layout esperado para esta empresa.",
                file_name=sanitized_file_name,
                file_size_bytes=file_size_bytes,
                detected_columns=detected_columns,
                missing_columns=best_alternative_probe.missing_columns,
                guidance=[
                    "Exporte o arquivo novamente usando o delimitador "
                    f"'{_display_csv_delimiter(tenant_config.csv.delimiter)}'.",
                    "O arquivo atual parece usar o delimitador "
                    f"'{_display_csv_delimiter(best_alternative_probe.delimiter)}'.",
                ],
                issues=[
                    UploadPreflightIssue(
                        code="delimiter_mismatch",
                        message=(
                            "Cabecalho encontrado, mas com separador diferente do configurado."
                        ),
                    )
                ],
            )

        if (
            best_alternative_probe.encoding != tenant_config.csv.encoding
            and best_alternative_probe.delimiter == tenant_config.csv.delimiter
        ):
            _raise_upload_preflight_error(
                "A codificacao do arquivo nao corresponde ao layout esperado para esta empresa.",
                file_name=sanitized_file_name,
                file_size_bytes=file_size_bytes,
                detected_columns=detected_columns,
                missing_columns=best_alternative_probe.missing_columns,
                guidance=[
                    f"Salve o CSV com codificacao {tenant_config.csv.encoding}.",
                    f"O arquivo atual parece estar em {best_alternative_probe.encoding}.",
                ],
                issues=[
                    UploadPreflightIssue(
                        code="encoding_mismatch",
                        message=(
                            "Cabecalho encontrado apenas com codificacao diferente da esperada."
                        ),
                    )
                ],
            )

        _raise_upload_preflight_error(
            "Nao foi possivel reconciliar o layout do CSV com o perfil desta empresa.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=detected_columns,
            missing_columns=best_alternative_probe.missing_columns,
            guidance=[
                "Revise o delimitador, a codificacao e o cabecalho antes de reenviar.",
                "Gere um novo CSV a partir do layout configurado para esta empresa.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="layout_mismatch",
                    message=(
                        "O cabecalho so foi parcialmente reconhecido com "
                        "outra combinacao de leitura."
                    ),
                )
            ],
        )

    if not expected_probe.columns or expected_probe.matched_column_count == 0:
        _raise_upload_preflight_error(
            "A primeira linha nao contem um cabecalho compativel com o layout esperado.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=expected_probe.columns,
            missing_columns=expected_columns,
            guidance=[
                "Confirme se a primeira linha traz o cabecalho do CSV.",
                "Revise o delimitador e gere novamente o arquivo antes do upload.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="missing_header",
                    message="Nenhuma coluna obrigatoria foi reconhecida na primeira linha.",
                )
            ],
        )

    if expected_probe.missing_columns:
        _raise_upload_preflight_error(
            "Faltam colunas obrigatorias no cabecalho do CSV.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=expected_probe.columns,
            missing_columns=expected_probe.missing_columns,
            guidance=[
                "Inclua as colunas faltantes no cabecalho antes de reenviar.",
                "Mantenha os nomes das colunas alinhados com o layout configurado para a empresa.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="missing_columns",
                    message=(
                        "O cabecalho foi lido, mas ainda nao contem todas "
                        "as colunas necessarias."
                    ),
                )
            ],
        )

    try:
        df = pd.read_csv(
            BytesIO(content),
            dtype=str,
            **_build_csv_read_kwargs(tenant_config),
        )
    except pd.errors.EmptyDataError:
        _raise_upload_preflight_error(
            "O arquivo enviado esta vazio.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=expected_probe.columns,
            missing_columns=[],
            guidance=[
                "Confirme se a planilha tem cabecalho na primeira linha.",
                "Inclua ao menos uma linha de dados antes de reenviar o arquivo.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="file_empty",
                    message="Nenhum dado foi encontrado abaixo do cabecalho.",
                )
            ],
        )
    except UnicodeDecodeError:
        _raise_upload_preflight_error(
            "A codificacao do arquivo nao corresponde ao layout esperado para esta empresa.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=expected_probe.columns,
            missing_columns=[],
            guidance=[
                f"Salve o CSV com codificacao {tenant_config.csv.encoding}.",
                "Reexporte o arquivo antes de tentar novamente.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="encoding_mismatch",
                    message=(
                        f"Falha ao ler o arquivo com codificacao "
                        f"{tenant_config.csv.encoding}."
                    ),
                )
            ],
        )
    except pd.errors.ParserError:
        _raise_upload_preflight_error(
            "Nao foi possivel ler o CSV com o layout configurado para esta empresa.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=expected_probe.columns,
            missing_columns=[],
            guidance=[
                "Revise delimitador, aspas e cabecalho antes de reenviar.",
                "Se possivel, gere novamente o arquivo CSV a partir da origem.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="csv_unreadable",
                    message="A leitura completa do CSV falhou antes da criacao do job.",
                )
            ],
        )

    if expected_probe.columns and len(expected_probe.columns) == len(df.columns):
        df.columns = expected_probe.columns

    if df.empty:
        _raise_upload_preflight_error(
            "O arquivo nao contem linhas de dados para validar.",
            file_name=sanitized_file_name,
            file_size_bytes=file_size_bytes,
            detected_columns=expected_probe.columns,
            missing_columns=[],
            guidance=[
                "Mantenha o cabecalho na primeira linha e preencha ao menos um item.",
                "Reexporte o CSV depois de conferir que ha registros abaixo do cabecalho.",
            ],
            issues=[
                UploadPreflightIssue(
                    code="missing_rows",
                    message="Nenhuma linha de dados foi encontrada abaixo do cabecalho.",
                )
            ],
        )

    return _build_preflight_result(
        file_name=sanitized_file_name,
        file_size_bytes=file_size_bytes,
        detected_columns=expected_probe.columns,
        missing_columns=[],
        guidance=[],
        issues=[],
    )


def _dataframe_to_csv_bytes(
    df: pd.DataFrame,
    tenant_config: TenantConfig,
) -> bytes:
    csv_content = df.to_csv(
        index=False,
        sep=tenant_config.csv.delimiter,
    )
    return csv_content.encode(tenant_config.csv.encoding)


def prepare_upload_content_for_job(
    *,
    file_name: str | None,
    content: bytes,
    tenant_config: TenantConfig,
) -> tuple[str, bytes]:
    sanitized_file_name = _sanitize_job_file_name(file_name)
    storage_file_name = _build_storage_csv_file_name(sanitized_file_name)
    file_extension = _get_upload_extension(sanitized_file_name)

    if file_extension == ".csv":
        return storage_file_name, content

    if file_extension == ".xlsx":
        raw_headers, df = _read_xlsx_content(content)
        if raw_headers and len(raw_headers) == len(df.columns):
            df.columns = raw_headers
        try:
            return storage_file_name, _dataframe_to_csv_bytes(df, tenant_config)
        except UnicodeEncodeError as exc:
            raise ValueError(
                "A planilha XLSX contem caracteres incompativeis com o layout CSV desta empresa."
            ) from exc

    raise ValueError(f"Unsupported upload format: {file_extension or 'unknown'}")


def build_tenant_upload_template_xlsx(
    tenant_config: TenantConfig,
) -> tuple[bytes, str]:
    expected_columns = _get_expected_upload_columns(tenant_config)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = UPLOAD_TEMPLATE_DATA_SHEET_NAME
    worksheet.freeze_panes = "A2"

    for column_index, column_name in enumerate(expected_columns, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=column_name)
        cell.font = Font(bold=True)
        worksheet.column_dimensions[cell.column_letter].width = max(len(column_name) + 4, 16)

    guide_sheet = workbook.create_sheet(title=UPLOAD_TEMPLATE_GUIDE_SHEET_NAME)
    guide_sheet["A1"] = "Como preparar a planilha"
    guide_sheet["A1"].font = Font(bold=True)
    guide_lines = [
        (
            f"1. Preencha a aba '{UPLOAD_TEMPLATE_DATA_SHEET_NAME}' e mantenha "
            "o cabecalho na primeira linha."
        ),
        "2. O sistema aceita o envio direto deste arquivo em XLSX.",
        (
            "3. Se preferir salvar em CSV, preserve o delimitador "
            f"'{_display_csv_delimiter(tenant_config.csv.delimiter)}' e a codificacao "
            f"{tenant_config.csv.encoding}."
        ),
        "4. Colunas esperadas: " + ", ".join(expected_columns),
    ]
    for row_index, line in enumerate(guide_lines, start=2):
        guide_sheet.cell(row=row_index, column=1, value=line)

    output = BytesIO()
    workbook.save(output)
    workbook.close()

    tenant_id = canonicalize_tenant_id(tenant_config.tenant_id)
    return output.getvalue(), f"{tenant_id}_modelo_validacao.xlsx"


def _read_raw_csv_headers(file_path: Path, tenant_config: TenantConfig) -> list[str]:
    with file_path.open("r", encoding=tenant_config.csv.encoding, newline="") as file:
        reader = csv.reader(file, delimiter=tenant_config.csv.delimiter)
        try:
            return next(reader)
        except StopIteration:
            return []


def _read_tenant_csv(file_path: Path, tenant_config: TenantConfig) -> pd.DataFrame:
    df = pd.read_csv(
        file_path,
        dtype=str,
        **_build_csv_read_kwargs(tenant_config),
    )
    raw_headers = _read_raw_csv_headers(file_path, tenant_config)
    if raw_headers and len(raw_headers) == len(df.columns):
        df.columns = raw_headers
    return df


def _read_tenant_csv_for_job(
    file_path: Path,
    tenant_config: TenantConfig,
    cancellation_checkpoint: CancellationCheckpoint,
) -> pd.DataFrame:
    cancellation_checkpoint()
    raw_headers = _read_raw_csv_headers(file_path, tenant_config)
    cancellation_checkpoint()

    reader = pd.read_csv(
        file_path,
        dtype=str,
        chunksize=CSV_READ_CHUNK_SIZE,
        **_build_csv_read_kwargs(tenant_config),
    )
    chunks: list[pd.DataFrame] = []
    for chunk in reader:
        cancellation_checkpoint()
        chunks.append(chunk)

    df = (
        pd.concat(chunks, ignore_index=True)
        if chunks
        else pd.DataFrame(columns=raw_headers)
    )
    if raw_headers and len(raw_headers) == len(df.columns):
        df.columns = raw_headers
    cancellation_checkpoint()
    return df


def _write_tenant_csv(
    df: pd.DataFrame,
    file_path: Path,
    tenant_config: TenantConfig,
) -> None:
    df.to_csv(
        file_path,
        index=False,
        sep=tenant_config.csv.delimiter,
        encoding=tenant_config.csv.encoding,
    )


def _dataframe_to_raw_rows(
    df: pd.DataFrame,
    cancellation_checkpoint: CancellationCheckpoint | None = None,
) -> list[dict[str, object]]:
    raw_rows: list[dict[str, object]] = []
    for offset, row in enumerate(df.itertuples(index=False, name=None)):
        if (
            cancellation_checkpoint is not None
            and offset % COOPERATIVE_CHECKPOINT_ROW_INTERVAL == 0
        ):
            cancellation_checkpoint()
        raw_rows.append(dict(zip(df.columns, row, strict=False)))

    if cancellation_checkpoint is not None:
        cancellation_checkpoint()
    return raw_rows


def _cleanup_temp_artifacts(paths: list[Path]) -> None:
    for path in paths:
        path.unlink(missing_ok=True)


class JobCancellationRequestedError(Exception):
    pass


class UploadPreflightError(ValueError):
    def __init__(self, detail: str, result: UploadPreflightResult) -> None:
        super().__init__(detail)
        self.detail = detail
        self.result = result


class OperationalExportKind(StrEnum):
    DUPLICATES = "duplicates"
    PROBLEM_GROUP = "problem_group"


@dataclass(frozen=True)
class DuplicateCsvResolution:
    kept_row_index: int
    current_kept_row_index: int
    deleted_row_indices: list[int]
    remaining_rows: int
    merged_columns: list[str]


@dataclass(frozen=True)
class ReviewFlagUpdate:
    row_index: int
    previous_status: str
    status: str
    review_flags: list[dict[str, object]]


def _copy_llm_parallel_shared_context(shared_context: dict) -> dict:
    return {
        LLM_FORCE_REFRESH_PARAM: resolve_force_refresh(
            shared_context.get(LLM_FORCE_REFRESH_PARAM)
        )
    }


def _extract_llm_audit_metadata(shared_context: dict) -> dict[str, Any] | None:
    return snapshot_llm_audit_metadata(shared_context)


def _validate_parallel_rule_batch(
    engine: ValidationEngine,
    *,
    batch_indices: list[int],
    normalized_rows: list[dict[str, str | int | float | None]],
    raw_rows: list[dict[str, object]],
    scoped_rows: list[dict[str, str | int | float | None]],
    rule_names: tuple[str, ...],
    base_shared_context: dict,
    shared_context: dict,
    max_workers: int,
    cancellation_checkpoint: CancellationCheckpoint,
) -> dict[int, list[ValidationIssue]]:
    if not batch_indices or not rule_names:
        return {}

    def _validate_row_with_isolated_context(
        row_index: int,
    ) -> tuple[list[ValidationIssue], dict[str, Any] | None]:
        row_shared_context = _copy_llm_parallel_shared_context(base_shared_context)
        issues = engine.validate_row(
            row_index=row_index,
            normalized_row=normalized_rows[row_index],
            raw_row=raw_rows[row_index],
            all_rows=scoped_rows,
            shared_context=row_shared_context,
            rule_names=rule_names,
        )
        return issues, _extract_llm_audit_metadata(row_shared_context)

    executor = ThreadPoolExecutor(max_workers=max_workers)
    future_to_row_index: dict[
        Future[tuple[list[ValidationIssue], dict[str, Any] | None]],
        int,
    ] = {
        executor.submit(_validate_row_with_isolated_context, row_index): row_index
        for row_index in batch_indices
    }
    pending = set(future_to_row_index)
    batch_results: dict[
        int,
        tuple[list[ValidationIssue], dict[str, Any] | None],
    ] = {}

    try:
        while pending:
            cancellation_checkpoint()
            done, pending = wait(
                pending,
                timeout=PARALLEL_CANCELLATION_POLL_SECONDS,
                return_when=FIRST_COMPLETED,
            )
            for future in done:
                row_index = future_to_row_index[future]
                batch_results[row_index] = future.result()
        cancellation_checkpoint()
    except BaseException:
        for future in pending:
            future.cancel()
        executor.shutdown(wait=False, cancel_futures=True)
        raise
    else:
        executor.shutdown(wait=True)

    results: dict[int, list[ValidationIssue]] = {}
    for row_index in batch_indices:
        issues, metadata = batch_results[row_index]
        merge_llm_audit_metadata(shared_context, metadata)
        results[row_index] = issues

    return results


def _get_job_csv_file(
    job_id: str,
    job_service: JobService,
) -> tuple[JobRecord, Path]:
    job = job_service.get_job(job_id)
    if job is None:
        raise KeyError(f"Job not found: {job_id}")

    if not job.file_path and not job.file_name:
        raise FileNotFoundError("Job does not contain a CSV file path")

    file_path = _resolve_job_upload_path(job)
    if not file_path.exists():
        raise FileNotFoundError("CSV file not found for this job")

    return job, file_path


def _get_job_csv_context(
    job_id: str,
    job_service: JobService,
) -> tuple[JobRecord, Path, TenantConfig, pd.DataFrame]:
    job, file_path = _get_job_csv_file(job_id, job_service)
    signature = _get_job_csv_signature(file_path)
    cached_context = _JOB_CSV_CONTEXT_CACHE.get(job_id)
    if (
        cached_context
        and cached_context.tenant_id == job.tenant_id
        and cached_context.file_path == file_path
        and cached_context.signature == signature
    ):
        return job, file_path, cached_context.tenant_config, cached_context.df

    tenant_config = load_tenant_config(job.tenant_id)
    df = _read_tenant_csv(file_path, tenant_config)
    _store_job_csv_context(job_id, job, file_path, tenant_config, df)
    return job, file_path, tenant_config, df


def _get_job_csv_signature(file_path: Path) -> tuple[int, int]:
    file_stat = file_path.stat()
    return file_stat.st_mtime_ns, file_stat.st_size


def _store_job_csv_context(
    job_id: str,
    job: JobRecord,
    file_path: Path,
    tenant_config: TenantConfig,
    df: pd.DataFrame,
) -> None:
    _JOB_CSV_CONTEXT_CACHE[job_id] = _CachedJobCsvContext(
        tenant_id=job.tenant_id,
        file_path=file_path,
        signature=_get_job_csv_signature(file_path),
        tenant_config=tenant_config,
        df=df,
    )


def _build_corrected_csv_name(job: JobRecord, file_path: Path) -> str:
    original_name = job.file_name or file_path.name
    stem = Path(original_name).stem or Path(file_path.name).stem or "lote"

    if stem.endswith("_corrigido"):
        return f"{stem}.csv"

    return f"{stem}_corrigido.csv"


def get_job_csv_download(
    job_id: str,
    job_service: JobService,
) -> tuple[Path, str]:
    job, file_path = _get_job_csv_file(job_id, job_service)
    return file_path, _build_corrected_csv_name(job, file_path)


def _get_job_result_data(
    job_id: str,
    job_service: JobService,
) -> tuple[JobRecord, dict]:
    job = job_service.get_job(job_id)
    if job is None:
        raise KeyError(f"Job not found: {job_id}")

    if job.status != JobStatus.COMPLETED:
        raise ValueError(f"Job not completed: {job.status.value}")

    result_path = _resolve_job_result_path(job)
    if not result_path.exists():
        raise FileNotFoundError("Result file not found")

    return job, json.loads(result_path.read_text(encoding="utf-8"))


def _coerce_row_index(value: object) -> int | None:
    if isinstance(value, bool):
        return None

    try:
        row_index = int(value)
    except (TypeError, ValueError):
        return None

    if row_index < 0:
        return None
    return row_index


def _collect_result_row_indices(payload: dict) -> set[int]:
    row_indices: set[int] = set()

    for row_result in payload.get("row_results") or []:
        if not isinstance(row_result, dict):
            continue
        row_index = _coerce_row_index(row_result.get("row_index"))
        if row_index is not None:
            row_indices.add(row_index)

    grouped_problems = payload.get("grouped_problems") or {}
    if isinstance(grouped_problems, dict):
        for occurrences in grouped_problems.values():
            if not isinstance(occurrences, list):
                continue
            for occurrence in occurrences:
                if not isinstance(occurrence, dict):
                    continue
                row_index = _coerce_row_index(occurrence.get("row_index"))
                if row_index is not None:
                    row_indices.add(row_index)

    for duplicate in payload.get("duplicates") or []:
        if not isinstance(duplicate, dict):
            continue
        for raw_row_index in duplicate.get("row_indices") or []:
            row_index = _coerce_row_index(raw_row_index)
            if row_index is not None:
                row_indices.add(row_index)

    return row_indices


def _load_review_flag_indices(job: JobRecord) -> set[int]:
    flags_path = _resolve_job_review_flags_path(job)
    if not flags_path.exists():
        return set()

    payload = json.loads(flags_path.read_text(encoding="utf-8"))
    raw_flags = payload.get("flags", {})
    flag_indices: set[int] = set()

    if isinstance(raw_flags, dict):
        iterable_flags = raw_flags.items()
    elif isinstance(raw_flags, list):
        iterable_flags = (
            (entry.get("row_index"), entry.get("status"))
            for entry in raw_flags
            if isinstance(entry, dict)
        )
    else:
        iterable_flags = ()

    for raw_row_index, raw_status in iterable_flags:
        row_index = _coerce_row_index(raw_row_index)
        if row_index is not None and raw_status == "review":
            flag_indices.add(row_index)

    return flag_indices


def _build_review_flags_payload(
    row_indices: set[int],
    *,
    allowed_row_indices: set[int] | None = None,
) -> list[dict[str, object]]:
    if allowed_row_indices is not None:
        row_indices = row_indices.intersection(allowed_row_indices)

    return [
        {"row_index": row_index, "status": "review"}
        for row_index in sorted(row_indices)
    ]


def _write_review_flag_indices(job: JobRecord, row_indices: set[int]) -> None:
    flags_path = _resolve_job_review_flags_path(job)
    flags_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "job_id": job.job_id,
        "tenant_id": job.tenant_id,
        "updated_at": datetime.now(UTC).isoformat(),
        "flags": {str(row_index): "review" for row_index in sorted(row_indices)},
    }
    flags_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _add_review_flags_to_result_payload(job: JobRecord, payload: dict) -> dict:
    result_payload = dict(payload)
    result_row_indices = _collect_result_row_indices(result_payload)
    result_payload["review_flags"] = _build_review_flags_payload(
        _load_review_flag_indices(job),
        allowed_row_indices=result_row_indices,
    )
    return result_payload


def get_job_result_payload(
    job_id: str,
    job_service: JobService,
) -> dict:
    job, payload = _get_job_result_data(job_id, job_service)
    return _add_review_flags_to_result_payload(job, payload)


def set_job_row_review_flag(
    job_id: str,
    job_service: JobService,
    *,
    row_index: int,
    status: str,
) -> ReviewFlagUpdate:
    if row_index < 0:
        raise ValueError(f"Invalid row index: {row_index}")
    if status not in {"review", "clear"}:
        raise ValueError("Review flag status must be 'review' or 'clear'")

    job, payload = _get_job_result_data(job_id, job_service)
    result_row_indices = _collect_result_row_indices(payload)
    if row_index not in result_row_indices:
        raise ValueError(f"Row index not found in job result: {row_index}")

    review_flag_indices = _load_review_flag_indices(job)
    previous_status = "review" if row_index in review_flag_indices else "clear"
    if status == "review":
        review_flag_indices.add(row_index)
    else:
        review_flag_indices.discard(row_index)

    _write_review_flag_indices(job, review_flag_indices)
    return ReviewFlagUpdate(
        row_index=row_index,
        previous_status=previous_status,
        status=status,
        review_flags=_build_review_flags_payload(
            review_flag_indices,
            allowed_row_indices=result_row_indices,
        ),
    )


def get_job_review_flags_path(
    job_id: str,
    job_service: JobService,
) -> Path:
    job = job_service.get_job(job_id)
    if job is None:
        raise KeyError(f"Job not found: {job_id}")
    return _resolve_job_review_flags_path(job)


def get_job_report_download(
    job_id: str,
    job_service: JobService,
) -> Path:
    job = job_service.get_job(job_id)
    if job is None:
        raise KeyError(f"Job not found: {job_id}")
    if job.status != JobStatus.COMPLETED:
        raise ValueError(f"Job not completed: {job.status.value}")

    report_path = _resolve_job_report_path(job)
    if not report_path.exists():
        raise FileNotFoundError("Report file not found")
    return report_path


def _sanitize_export_token(value: str) -> str:
    sanitized = re.sub(r"[^a-zA-Z0-9]+", "_", value).strip("_").lower()
    return sanitized or "grupo"


def _build_operational_export_name(
    job: JobRecord,
    export_kind: OperationalExportKind,
    problem_code: str | None = None,
) -> str:
    original_name = job.file_name or "lote.csv"
    original_path = Path(original_name)
    stem = original_path.stem or "lote"
    suffix = original_path.suffix or ".csv"

    if export_kind == OperationalExportKind.DUPLICATES:
        return f"{stem}_duplicados{suffix}"

    if problem_code is None:
        raise ValueError("problem_code is required for problem_group exports")

    return f"{stem}_{_sanitize_export_token(problem_code)}{suffix}"


def get_job_operational_export(
    job_id: str,
    job_service: JobService,
    *,
    export_kind: OperationalExportKind,
    problem_code: str | None = None,
) -> tuple[str, str]:
    job, report_data = _get_job_result_data(job_id, job_service)

    if export_kind == OperationalExportKind.DUPLICATES:
        duplicates = report_data.get("duplicates") or []
        if not duplicates:
            raise KeyError("No duplicate rows available for export")
        return (
            build_duplicates_export_csv(duplicates),
            _build_operational_export_name(job, export_kind),
        )

    if problem_code is None:
        raise ValueError("problem_code is required for problem_group exports")

    grouped_problems = report_data.get("grouped_problems") or {}
    occurrences = grouped_problems.get(problem_code)
    if not occurrences:
        raise KeyError(f"Problem group not found: {problem_code}")

    return (
        build_problem_group_export_csv(problem_code, occurrences),
        _build_operational_export_name(job, export_kind, problem_code),
    )


def _resolve_source_column(
    field_name: str,
    tenant_columns: dict[str, str],
    available_columns: list[str],
) -> str:
    candidate_columns = [
        tenant_columns.get(field_name),
        DEFAULT_TENANT_COLUMNS.get(field_name),
        field_name,
    ]

    for candidate in candidate_columns:
        if not candidate:
            continue
        resolved = resolve_source_column_name(candidate, available_columns)
        if resolved is not None:
            return resolved

    raise ValueError(f"Column not found for field: {field_name}")


def _build_canonical_to_source_column_map(
    tenant_config: TenantConfig,
    available_columns: list[str],
) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for canonical_field in DEFAULT_TENANT_COLUMNS:
        try:
            mapping[canonical_field] = _resolve_source_column(
                canonical_field, tenant_config.columns, available_columns
            )
        except ValueError:
            continue
    return mapping


def _build_xlsx_export_name(job: JobRecord, file_path: Path) -> str:
    original_name = job.file_name or file_path.name
    original_path = Path(original_name)
    stem = original_path.stem or "lote"
    if stem.endswith("_corrigido"):
        return f"{stem}.xlsx"
    return f"{stem}_corrigido.xlsx"


def get_job_xlsx_export(
    job_id: str,
    job_service: JobService,
) -> tuple[bytes, str]:
    job, file_path = _get_job_csv_file(job_id, job_service)
    _, report_data = _get_job_result_data(job_id, job_service)

    _, _, tenant_config, df = _get_job_csv_context(job_id, job_service)
    source_columns = df.columns.tolist()
    source_rows = _dataframe_to_raw_rows(df)
    row_results = report_data.get("row_results") or []
    canonical_to_source_column = _build_canonical_to_source_column_map(
        tenant_config, source_columns
    )

    xlsx_bytes = build_xlsx_export_bytes(
        source_rows=source_rows,
        source_columns=source_columns,
        row_results=row_results,
        canonical_to_source_column=canonical_to_source_column,
    )
    return xlsx_bytes, _build_xlsx_export_name(job, file_path)


def update_job_csv_row(
    job_id: str,
    job_service: JobService,
    *,
    row_index: int,
    updates: dict[str, str],
) -> dict[str, str]:
    if row_index < 0:
        raise ValueError(f"Invalid row index: {row_index}")

    job, file_path, tenant_config, df = _get_job_csv_context(job_id, job_service)
    if row_index >= len(df):
        raise ValueError(f"Row index out of range: {row_index}")

    for field_name, new_value in updates.items():
        column_name = _resolve_source_column(
            field_name, tenant_config.columns, df.columns.tolist()
        )
        df.at[row_index, column_name] = new_value

    _write_tenant_csv(df, file_path, tenant_config)
    _store_job_csv_context(job_id, job, file_path, tenant_config, df)
    return df.iloc[row_index].to_dict()  # type: ignore[return-value]


def delete_job_csv_rows(
    job_id: str,
    job_service: JobService,
    *,
    row_indices: list[int],
) -> int:
    normalized_indices = sorted(set(row_indices))
    if not normalized_indices:
        raise ValueError("At least one row index must be provided")
    if normalized_indices[0] < 0:
        raise ValueError(f"Invalid row index: {normalized_indices[0]}")

    job, file_path, tenant_config, df = _get_job_csv_context(job_id, job_service)
    if normalized_indices[-1] >= len(df):
        raise ValueError(f"Row index out of range: {normalized_indices[-1]}")

    df = df.drop(index=normalized_indices).reset_index(drop=True)
    _write_tenant_csv(df, file_path, tenant_config)
    _store_job_csv_context(job_id, job, file_path, tenant_config, df)
    return len(df)


def replace_job_csv_rows(
    job_id: str,
    job_service: JobService,
    *,
    rows: list[dict[str, object]],
) -> int:
    job, file_path, tenant_config, df = _get_job_csv_context(job_id, job_service)
    columns = [str(column) for column in df.columns.tolist()]
    normalized_rows = [
        {column: row.get(column, "") for column in columns}
        for row in rows
    ]
    updated_df = pd.DataFrame(normalized_rows, columns=columns)
    _write_tenant_csv(updated_df, file_path, tenant_config)
    _store_job_csv_context(job_id, job, file_path, tenant_config, updated_df)
    return len(updated_df)


def _is_missing_csv_value(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()

    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _normalize_csv_column_name(column_name: object) -> str:
    text = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "_", str(column_name or ""))
    normalized = unicodedata.normalize("NFD", text.strip().lower())
    normalized = "".join(
        char for char in normalized if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"[^a-z0-9]+", "_", normalized).strip("_")


def _get_normalized_csv_column_tokens(column_name: object) -> set[str]:
    normalized_name = _normalize_csv_column_name(column_name)
    if not normalized_name:
        return set()
    return {token for token in normalized_name.split("_") if token}


def _is_media_like_column(column_name: object) -> bool:
    normalized_name = _normalize_csv_column_name(column_name)
    if not normalized_name:
        return False
    if normalized_name in MEDIA_COLUMN_NAMES:
        return True

    tokens = _get_normalized_csv_column_tokens(column_name)
    return any(token in MEDIA_COLUMN_TOKENS for token in tokens)


def _looks_like_date_or_time_value(value: object) -> bool:
    if _is_missing_csv_value(value):
        return False

    text = str(value).strip()
    if not text:
        return False

    return any(pattern.fullmatch(text) for pattern in DATE_TIME_VALUE_PATTERNS)


def _is_date_or_time_like_column(
    column_name: object,
    sample_values: list[object],
) -> bool:
    normalized_name = _normalize_csv_column_name(column_name)
    if not normalized_name:
        return False

    if normalized_name.endswith("_at"):
        return True

    tokens = _get_normalized_csv_column_tokens(column_name)
    if any(token in DATE_TIME_COLUMN_TOKENS for token in tokens):
        return True

    non_empty_values = [value for value in sample_values if not _is_missing_csv_value(value)]
    if not non_empty_values:
        return False

    inspected_values = non_empty_values[:5]
    return len(inspected_values) >= 2 and all(
        _looks_like_date_or_time_value(value) for value in inspected_values
    )


def _get_same_name_merge_skipped_columns(
    df: pd.DataFrame,
    row_indices: list[int],
) -> set[str]:
    skipped_columns: set[str] = set()

    for column_name in df.columns:
        sample_values = [df.at[row_index, column_name] for row_index in row_indices]
        if _is_media_like_column(column_name) or _is_date_or_time_like_column(
            column_name,
            sample_values,
        ):
            skipped_columns.add(str(column_name))

    return skipped_columns


def _has_duplicate_description_conflict(
    df: pd.DataFrame,
    tenant_config: TenantConfig,
    row_indices: list[int],
) -> bool:
    try:
        description_column = _resolve_source_column(
            "descricao",
            tenant_config.columns,
            df.columns.tolist(),
        )
    except ValueError:
        return True

    description_values: list[str] = []
    seen_values: set[str] = set()
    for row_index in row_indices:
        description_value = df.at[row_index, description_column]
        if _is_missing_csv_value(description_value):
            continue

        description_text = str(description_value).strip()
        if not description_text or description_text in seen_values:
            continue

        seen_values.add(description_text)
        description_values.append(description_text)
        if len(description_values) > 1:
            return True

    return False


def _merge_missing_duplicate_values(
    df: pd.DataFrame,
    *,
    keep_row_index: int,
    source_row_indices: list[int],
    skipped_columns: set[str] | None = None,
) -> list[str]:
    merged_columns: list[str] = []
    source_indices = sorted(set(source_row_indices), reverse=True)
    skipped = skipped_columns or set()

    for column_name in df.columns:
        if str(column_name) in skipped:
            continue

        current_value = df.at[keep_row_index, column_name]
        if not _is_missing_csv_value(current_value):
            continue

        for source_index in source_indices:
            source_value = df.at[source_index, column_name]
            if _is_missing_csv_value(source_value):
                continue

            df.at[keep_row_index, column_name] = source_value
            merged_columns.append(str(column_name))
            break

    return merged_columns


def rerun_job_validation(
    job_id: str,
    job_service: JobService,
) -> JobRecord:
    job = job_service.get_job(job_id)
    if job is None:
        raise KeyError(f"Job not found: {job_id}")
    if job.status != JobStatus.COMPLETED:
        raise ValueError("Only completed jobs can be revalidated in place")

    run_validation_job(job_id, job_service)

    refreshed_job = job_service.get_job(job_id)
    if refreshed_job is None:
        raise KeyError(f"Job not found: {job_id}")
    if refreshed_job.status != JobStatus.COMPLETED:
        raise RuntimeError(
            refreshed_job.error_message
            or "Failed to rebuild the current job artifacts"
        )
    return refreshed_job


def delete_job_csv_rows_and_refresh(
    job_id: str,
    job_service: JobService,
    *,
    row_indices: list[int],
) -> int:
    job = job_service.get_job(job_id)
    if job is None:
        raise KeyError(f"Job not found: {job_id}")
    if job.status != JobStatus.COMPLETED:
        raise ValueError("Only completed jobs can delete rows and refresh the job")

    remaining_rows = delete_job_csv_rows(
        job_id,
        job_service,
        row_indices=row_indices,
    )
    rerun_job_validation(job_id, job_service)
    return remaining_rows


def resolve_duplicate_csv_rows_and_refresh(
    job_id: str,
    job_service: JobService,
    *,
    row_indices: list[int],
) -> DuplicateCsvResolution:
    normalized_indices = sorted(set(row_indices))
    if not normalized_indices:
        raise ValueError("At least one row index must be provided")
    if normalized_indices[0] < 0:
        raise ValueError(f"Invalid row index: {normalized_indices[0]}")

    job = job_service.get_job(job_id)
    if job is None:
        raise KeyError(f"Job not found: {job_id}")
    if job.status != JobStatus.COMPLETED:
        raise ValueError("Only completed jobs can resolve duplicates and refresh the job")

    keep_row_index = normalized_indices[-1]
    deleted_row_indices = [
        row_index for row_index in normalized_indices if row_index != keep_row_index
    ]
    if not deleted_row_indices:
        raise ValueError("At least one duplicate row must be removed")

    job, file_path, tenant_config, df = _get_job_csv_context(job_id, job_service)
    if normalized_indices[-1] >= len(df):
        raise ValueError(f"Row index out of range: {normalized_indices[-1]}")

    has_description_conflict = _has_duplicate_description_conflict(
        df,
        tenant_config,
        normalized_indices,
    )
    skipped_columns: set[str] | None = None
    if not has_description_conflict:
        skipped_columns = _get_same_name_merge_skipped_columns(df, normalized_indices)

    merged_columns = _merge_missing_duplicate_values(
        df,
        keep_row_index=keep_row_index,
        source_row_indices=deleted_row_indices,
        skipped_columns=skipped_columns,
    )

    df = df.drop(index=deleted_row_indices).reset_index(drop=True)
    _write_tenant_csv(df, file_path, tenant_config)
    _store_job_csv_context(job_id, job, file_path, tenant_config, df)

    rerun_job_validation(job_id, job_service)
    current_kept_row_index = keep_row_index - sum(
        1 for row_index in deleted_row_indices if row_index < keep_row_index
    )
    return DuplicateCsvResolution(
        kept_row_index=keep_row_index,
        current_kept_row_index=current_kept_row_index,
        deleted_row_indices=deleted_row_indices,
        remaining_rows=len(df),
        merged_columns=merged_columns,
    )


def read_job_csv_row(
    job_id: str,
    job_service: JobService,
    *,
    row_index: int,
) -> tuple[dict[str, str], dict[str, str]]:
    if row_index < 0:
        raise ValueError(f"Invalid row index: {row_index}")

    _, _, tenant_config, df = _get_job_csv_context(job_id, job_service)
    if row_index >= len(df):
        raise ValueError(f"Row index out of range: {row_index}")

    row = {
        str(column): "" if value is None else str(value)
        for column, value in df.iloc[row_index].to_dict().items()
    }
    resolved_columns = {
        field_name: tenant_config.columns.get(field_name, field_name)
        for field_name in tenant_config.columns
    }
    return row, resolved_columns


def read_all_job_csv_rows(
    job_id: str,
    job_service: JobService,
) -> list[dict[str, str]]:
    _, _, _, df = _get_job_csv_context(job_id, job_service)
    return [
        {str(column): "" if value is None else str(value) for column, value in row.items()}
        for row in df.to_dict(orient="records")
    ]


def create_reprocess_job(
    job_id: str,
    job_service: JobService,
    *,
    api_key_id: str | None = None,
    force_refresh: bool | None = None,
) -> JobRecord:
    source_job, source_file_path, _, _ = _get_job_csv_context(job_id, job_service)
    source_file_name = source_job.file_name or source_file_path.name

    params = dict(source_job.params)
    if force_refresh is not None:
        params[LLM_FORCE_REFRESH_PARAM] = force_refresh

    new_job = job_service.create_job(
        tenant_id=source_job.tenant_id,
        file_name=source_file_name,
        api_key_id=api_key_id,
        params=params,
        parent_job_id=source_job.job_id,
    )
    destination_path = build_job_upload_path(
        new_job.tenant_id,
        new_job.job_id,
        _build_storage_csv_file_name(source_file_name),
    )
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source_file_path, destination_path)

    new_job.file_path = str(destination_path)
    new_job.file_name = source_file_name
    job_service.save_job(new_job.job_id)
    job_service.register_reprocess_link(
        source_job_id=source_job.job_id,
        new_job_id=new_job.job_id,
    )
    return new_job


def ensure_rules_registered() -> None:
    from app.core.registry import RULE_REGISTRY

    if RULE_REGISTRY:
        return
    register_rule(DuplicateItemRule())
    register_rule(FlagConsistencyRule())
    register_rule(BrandModelConsistencyRule())
    register_rule(ZeroItemQualityRule())
    register_rule(CategoryRequiredFieldsRule())
    register_rule(CategoryCriticalCheckRule())
    register_rule(SuspiciousPatternRule())
    register_rule(LLMAuditRule())


def _resolve_batch_size(total_rows: int, *, parallel_workers: int = 1) -> int:
    if total_rows <= 0:
        return 1

    target_size = math.ceil(total_rows / TARGET_PREVIEW_BATCHES)
    target_size = max(target_size, parallel_workers)
    return max(1, min(DEFAULT_BATCH_SIZE, target_size))


def _build_indexing_status_detail(
    validated_rows: int,
    source_total_rows: int,
    validation_scope: ValidationScope,
) -> str:
    if validated_rows == 0:
        if validation_scope == ValidationScope.ALL_ITEMS:
            return (
                "A indexação global foi concluída. O arquivo foi lido, mas não "
                "há linhas de dados para validar neste lote."
            )
        if validation_scope == ValidationScope.DUPLICATE_ITEMS:
            return (
                "A indexação global foi concluída. O arquivo foi lido, mas nenhuma "
                "linha entrou no escopo operacional porque não há grupos de "
                "duplicidade no lote."
            )
        return (
            "A indexação global foi concluída. O arquivo foi lido, mas nenhuma linha "
            "entrou no escopo operacional porque não há itens cadastrados do zero."
        )

    if validation_scope == ValidationScope.ALL_ITEMS:
        return (
            "A indexação global foi concluída. O arquivo tem "
            f"{source_total_rows} linhas e todas entrarão na validação operacional "
            "e na prévia."
        )

    if validation_scope == ValidationScope.DUPLICATE_ITEMS:
        return (
            "A indexação global foi concluída. O arquivo tem "
            f"{source_total_rows} linhas e {validated_rows} linhas pertencentes a "
            "grupos com Item duplicado entrarão na validação operacional e na prévia."
        )

    return (
        "A indexação global foi concluída. O arquivo tem "
        f"{source_total_rows} linhas e {validated_rows} itens cadastrados do zero "
        "entrarão na validação operacional e na prévia."
    )


def _build_validation_status_detail(
    processed_rows: int,
    validated_rows: int,
    source_total_rows: int,
) -> str:
    return (
        "A prévia operacional está sendo atualizada em batches seguros após a "
        "indexação global do lote. "
        f"{processed_rows} de {validated_rows} itens em escopo já foram validados "
        f"com contexto completo. O CSV original tem {source_total_rows} linhas."
    )


def _get_job_validation_scope(job: JobRecord) -> ValidationScope:
    return parse_validation_scope(job.params.get(VALIDATION_SCOPE_PARAM))


def _build_scope_indexing_detail(validation_scope: ValidationScope) -> str:
    if validation_scope == ValidationScope.ALL_ITEMS:
        return (
            "O sistema está mapeando o conjunto completo do arquivo para "
            "confirmar que todas as linhas do lote entram no resultado "
            "operacional e liberar apenas prévias compatíveis com esse escopo."
        )

    if validation_scope == ValidationScope.DUPLICATE_ITEMS:
        return (
            "O sistema está mapeando o conjunto completo do arquivo para "
            "identificar quais linhas pertencem a grupos com Item duplicado e "
            "liberar apenas prévias compatíveis com esse escopo."
        )

    return (
        "O sistema está mapeando o conjunto completo do arquivo para "
        "identificar quais itens cadastrados do zero entram no resultado "
        "operacional e liberar apenas prévias compatíveis com esse escopo."
    )


def _build_scope_artifact_detail(validation_scope: ValidationScope) -> str:
    if validation_scope == ValidationScope.ALL_ITEMS:
        return (
            "O sistema está montando o resumo executivo, os dados estruturados "
            "e o relatório PDF com todas as linhas que entraram no escopo "
            "validado deste lote."
        )

    if validation_scope == ValidationScope.DUPLICATE_ITEMS:
        return (
            "O sistema está montando o resumo executivo, os dados estruturados "
            "e o relatório PDF apenas com as linhas que pertencem a grupos com "
            "Item duplicado e entraram no escopo validado."
        )

    return (
        "O sistema está montando o resumo executivo, os dados estruturados "
        "e o relatório PDF apenas com os itens cadastrados do zero que "
        "entraram no escopo validado."
    )


def _raise_if_cancellation_requested(
    job_id: str,
    job_service: JobService,
) -> None:
    job = job_service.get_job(job_id)
    if job is None:
        raise JobCancellationRequestedError(
            "O lote foi interrompido porque o job deixou de existir."
        )
    if job.status == JobStatus.CANCELED or job.cancel_requested:
        raise JobCancellationRequestedError(
            "O processamento foi cancelado por solicitação do usuário."
        )


def run_validation_job(
    job_id: str,
    job_service: JobService,
    *,
    worker_id: str | None = None,
    execution_lease_seconds: int = DEFAULT_EXECUTION_LEASE_SECONDS,
) -> None:
    job = job_service.get_job(job_id)
    if job is None or job.status == JobStatus.CANCELED:
        return

    started_at = datetime.now(UTC)
    temp_artifact_paths: list[Path] = []
    last_execution_heartbeat = 0.0
    heartbeat_interval_seconds = max(
        1.0,
        min(float(execution_lease_seconds) / 3.0, 30.0),
    )

    def maybe_refresh_execution_lease(*, force: bool = False) -> None:
        nonlocal last_execution_heartbeat
        if worker_id is None:
            return
        now_monotonic = time.monotonic()
        if not force and (
            now_monotonic - last_execution_heartbeat
        ) < heartbeat_interval_seconds:
            return
        refreshed_job = job_service.heartbeat_job_execution(
            job_id,
            worker_id=worker_id,
            lease_seconds=execution_lease_seconds,
        )
        if refreshed_job is None:
            raise RuntimeError(
                "The worker lost ownership of the execution lease for this job."
            )
        last_execution_heartbeat = now_monotonic

    def cancellation_checkpoint() -> None:
        maybe_refresh_execution_lease()
        _raise_if_cancellation_requested(job_id, job_service)

    log_event(
        _logger,
        "validation.started",
        tenant_id=job.tenant_id,
        job_id=job_id,
    )

    try:
        cancellation_checkpoint()
        job = job_service.begin_job_execution(job_id)
        maybe_refresh_execution_lease(force=True)
        cancellation_checkpoint()

        ensure_rules_registered()

        tenant_config = load_tenant_config(job.tenant_id)
        engine = ValidationEngine(tenant=tenant_config)
        validation_scope = _get_job_validation_scope(job)

        file_path = Path(job.file_path)  # type: ignore[arg-type]
        df = _read_tenant_csv_for_job(
            file_path,
            tenant_config,
            cancellation_checkpoint,
        )
        raw_rows = _dataframe_to_raw_rows(
            df,
            cancellation_checkpoint=cancellation_checkpoint,
        )

        job_service.update_progress(
            job_id=job_id,
            current_step="indexing_global",
            status_title="Indexação global do lote",
            status_detail=_build_scope_indexing_detail(validation_scope),
        )
        cancellation_checkpoint()
        normalized_rows = engine.normalize_rows(
            raw_rows,
            cancellation_checkpoint=cancellation_checkpoint,
        )
        source_total_rows = len(normalized_rows)
        scoped_row_indices = engine.get_scoped_row_indices(
            normalized_rows,
            validation_scope=validation_scope,
            cancellation_checkpoint=cancellation_checkpoint,
        )
        scoped_rows = [normalized_rows[idx] for idx in scoped_row_indices]
        total_rows = len(scoped_row_indices)
        execution_plan: ValidationExecutionPlan = engine.build_execution_plan(
            scoped_row_count=total_rows
        )
        batch_size = _resolve_batch_size(
            total_rows,
            parallel_workers=execution_plan.parallel_workers,
        )
        partial_duplicates = build_duplicate_section(
            normalized_rows,
            {},
            row_indices=scoped_row_indices,
        )

        validation_results: dict[int, list[ValidationIssue]] = {}
        processed_row_indices: list[int] = []
        shared_context: dict = {
            LLM_FORCE_REFRESH_PARAM: resolve_force_refresh(
                job.params.get(LLM_FORCE_REFRESH_PARAM)
            )
        }

        job_service.update_partial_result(
            job_id=job_id,
            total_rows=total_rows,
            source_total_rows=source_total_rows,
            processed_rows=0,
            batch_size=batch_size,
            partial_summary={
                "total_rows": total_rows,
                "validated_rows": total_rows,
                "source_total_rows": source_total_rows,
                "processed_rows": 0,
                "rows_with_issues": 0,
                "total_issues": 0,
                "error_count": 0,
                "warning_count": 0,
            },
            partial_grouped_problems={},
            partial_duplicates=partial_duplicates,
            row_results_preview=[],
            is_partial_result_available=bool(partial_duplicates),
            current_step="validating_batches",
            status_title="Validação em batches iniciada",
            status_detail=_build_indexing_status_detail(
                total_rows,
                source_total_rows,
                validation_scope,
            ),
        )
        cancellation_checkpoint()

        for start in range(0, total_rows, batch_size):
            cancellation_checkpoint()
            stop = min(start + batch_size, total_rows)
            batch_indices = scoped_row_indices[start:stop]

            for idx in batch_indices:
                cancellation_checkpoint()
                validation_results[idx] = []
                if execution_plan.serial_rule_names:
                    validation_results[idx] = engine.validate_row(
                        row_index=idx,
                        normalized_row=normalized_rows[idx],
                        raw_row=raw_rows[idx],
                        all_rows=scoped_rows,
                        shared_context=shared_context,
                        rule_names=execution_plan.serial_rule_names,
                        cancellation_checkpoint=cancellation_checkpoint,
                    )
                cancellation_checkpoint()

            if execution_plan.parallel_rule_names:
                parallel_results = _validate_parallel_rule_batch(
                    engine,
                    batch_indices=batch_indices,
                    normalized_rows=normalized_rows,
                    raw_rows=raw_rows,
                    scoped_rows=scoped_rows,
                    rule_names=execution_plan.parallel_rule_names,
                    base_shared_context=shared_context,
                    shared_context=shared_context,
                    max_workers=execution_plan.parallel_workers,
                    cancellation_checkpoint=cancellation_checkpoint,
                )
                for idx in batch_indices:
                    validation_results[idx].extend(parallel_results.get(idx, []))
                cancellation_checkpoint()

            processed_row_indices.extend(batch_indices)
            partial_report = build_partial_report(
                normalized_rows,
                validation_results,
                processed_row_indices,
                partial_duplicates=partial_duplicates,
                validated_row_indices=scoped_row_indices,
                source_total_rows=source_total_rows,
                include_row_results_preview=False,
            )

            job_service.update_partial_result(
                job_id=job_id,
                total_rows=total_rows,
                source_total_rows=source_total_rows,
                processed_rows=len(processed_row_indices),
                batch_size=batch_size,
                partial_summary=partial_report["partial_summary"],
                partial_grouped_problems=partial_report["partial_grouped_problems"],
                partial_duplicates=partial_report["partial_duplicates"],
                row_results_preview=partial_report["row_results_preview"],
                is_partial_result_available=bool(
                    len(processed_row_indices) > 0
                    or partial_report["partial_grouped_problems"]
                    or partial_report["partial_duplicates"]
                ),
                current_step="validating_batches",
                status_title="Prévia operacional em atualização",
                status_detail=_build_validation_status_detail(
                    len(processed_row_indices),
                    total_rows,
                    source_total_rows,
                ),
            )
            cancellation_checkpoint()

        report_data = build_full_report(
            normalized_rows,
            validation_results,
            validated_row_indices=scoped_row_indices,
            source_total_rows=source_total_rows,
            llm_audit_metadata=shared_context.get(LLM_AUDIT_METADATA_KEY),
        )
        cancellation_checkpoint()

        job_service.update_progress(
            job_id=job_id,
            current_step="building_artifacts",
            status_title="Consolidação dos artefatos",
            status_detail=_build_scope_artifact_detail(validation_scope),
        )
        cancellation_checkpoint()

        result_path = build_job_result_path(job.tenant_id, job_id)
        result_path.parent.mkdir(parents=True, exist_ok=True)
        result_temp_path = result_path.with_name(f"{result_path.name}.tmp")
        temp_artifact_paths.append(result_temp_path)
        result_temp_path.write_text(
            json.dumps(report_data, ensure_ascii=False, default=str),
            encoding="utf-8",
        )
        cancellation_checkpoint()

        pdf_path = build_job_report_path(job.tenant_id, job_id)
        pdf_temp_path = pdf_path.with_name(f"{pdf_path.name}.tmp")
        temp_artifact_paths.append(pdf_temp_path)
        generate_pdf_report(
            normalized_rows,
            validation_results,
            pdf_temp_path,
            metadata={
                "organization_name": tenant_config.display_name,
                "file_name": job.file_name,
                "job_id": job_id,
                "generated_at": datetime.now(UTC),
            },
            validated_row_indices=scoped_row_indices,
            source_total_rows=source_total_rows,
            validation_scope=validation_scope,
            llm_audit_metadata=shared_context.get(LLM_AUDIT_METADATA_KEY),
        )
        cancellation_checkpoint()

        result_temp_path.replace(result_path)
        pdf_temp_path.replace(pdf_path)
        temp_artifact_paths.clear()

        summary = report_data["summary"]
        job_service.complete_job(
            job_id=job_id,
            result_path=str(result_path),
            report_path=str(pdf_path),
            total_rows=summary["total_rows"],
            source_total_rows=summary["source_total_rows"],
            rows_with_issues=summary["rows_with_issues"],
            total_issues=summary["total_issues"],
        )
        log_event(
            _logger,
            "validation.completed",
            tenant_id=job.tenant_id,
            job_id=job_id,
            duration_ms=(datetime.now(UTC) - started_at).total_seconds() * 1000.0,
            total_rows=summary["total_rows"],
        )

    except JobCancellationRequestedError as exc:
        _cleanup_temp_artifacts(temp_artifact_paths)
        current_job = job_service.get_job(job_id)
        if current_job is not None and current_job.status != JobStatus.CANCELED:
            job_service.cancel_job(job_id, str(exc))
        log_event(
            _logger,
            "validation.canceled",
            level="warning",
            tenant_id=job.tenant_id,
            job_id=job_id,
            duration_ms=(datetime.now(UTC) - started_at).total_seconds() * 1000.0,
            detail=str(exc),
        )
    except Exception as exc:
        _cleanup_temp_artifacts(temp_artifact_paths)
        current_job = job_service.get_job(job_id)
        if current_job is not None and current_job.status != JobStatus.CANCELED:
            job_service.fail_job(job_id, str(exc))
        log_event(
            _logger,
            "validation.failed",
            level="error",
            tenant_id=job.tenant_id,
            job_id=job_id,
            duration_ms=(datetime.now(UTC) - started_at).total_seconds() * 1000.0,
            error=str(exc),
        )
